"""稿件读取器：md / txt / docx（python-docx）。PDF 待预装 PyMuPDF。"""
from __future__ import annotations
from pathlib import Path


def read_document(path: str | Path) -> tuple[str, str]:
    """返回 (原始文本, 类型)。"""
    p = Path(path)
    suffix = p.suffix.lower()
    if suffix == ".docx":
        return _read_docx(p), "docx"
    if suffix == ".pdf":
        return _read_pdf(p), "pdf"
    return p.read_text(encoding="utf-8"), "text"


def _read_docx(p: Path) -> str:
    import docx  # 已安装
    d = docx.Document(str(p))
    parts = []
    for para in d.paragraphs:
        if para.text.strip():
            parts.append(para.text)
    for table in d.tables:
        for row in table.rows:
            cells = [c.text.strip() for c in row.cells]
            parts.append("| " + " | ".join(cells) + " |")
    return "\n".join(parts)


def _read_pdf(p: Path) -> str:
    # PyMuPDF (fitz) 未安装时的优雅降级
    try:
        import fitz  # noqa: F401
    except ImportError as e:
        raise RuntimeError(
            "PDF 解析需要 PyMuPDF。安装：pip install pymupdf（或 pip install docling 获得更强解析）。"
        ) from e
    doc = fitz.open(str(p))
    parts = [page.get_text() for page in doc]
    doc.close()
    return "\n".join(parts)
