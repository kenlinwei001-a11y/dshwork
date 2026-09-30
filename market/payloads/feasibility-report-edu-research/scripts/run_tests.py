#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
测试套件。零外部依赖，只用标准库 + PyYAML。

用法：
    python3 scripts/run_tests.py           跑全部
    python3 scripts/run_tests.py -v        显示每条断言

两种模式，同一份代码：
    base 模式    存在 references/verticals/<id>/ → 遍历全部垂类
    垂类模式     存在 references/vertical/       → 只跑这一个

**测试基线不硬编码在代码里，从各垂类的 meta.yaml 读。**
加一个垂类只需加一个目录，不需要改这个文件——
这跟本体层"事实只存一份"是同一个原则，只是对象从报告数字换成了测试期望。

三类测试：
  黄金用例  真实报告夹具，锁定检查结果。检查器改坏了这里会红
  故障注入  每种缺陷各造一个最小用例，确认对应检查项确实会报
  单元测试  表达式求值、环检测、引用提取、schema、工具脚本
"""

import json
import re
import sys
import tempfile
from collections import Counter
from pathlib import Path

import yaml

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT / "scripts"))

from schema import SchemaValidator, detect_cycles, _refs_in_formula  # noqa: E402
from assets import load_l0  # noqa: E402
import validators as V  # noqa: E402

VERBOSE = "-v" in sys.argv
_passed, _failed = 0, []


def check(name, cond, detail=""):
    global _passed
    if cond:
        _passed += 1
        if VERBOSE:
            print(f"  ok   {name}")
    else:
        _failed.append((name, detail))
        print(f"  FAIL {name}" + (f"\n       {detail}" if detail else ""))


def load_norms():
    p = ROOT / "references/norms.yaml"
    return yaml.safe_load(p.read_text(encoding="utf-8")) if p.exists() else {}


def verticals():
    """返回 [(目录, meta)]。垂类模式只有一个，base 模式有多个。"""
    single = ROOT / "references/vertical"
    if (single / "meta.yaml").exists():
        return [(single, yaml.safe_load((single / "meta.yaml").read_text(encoding="utf-8")))]
    out = []
    root = ROOT / "references/verticals"
    if root.exists():
        for d in sorted(root.iterdir()):
            m = d / "meta.yaml"
            if d.is_dir() and m.exists():
                out.append((d, yaml.safe_load(m.read_text(encoding="utf-8"))))
    return out


def run_all(proj_yaml_text):
    core, argu, deriv = load_l0()
    proj = yaml.safe_load(proj_yaml_text)
    sv = SchemaValidator(proj, core, argu)
    sf = sv.run()
    if sv.fatal:
        return sf, [], None
    v = V.Validator(core, argu, deriv, proj, load_norms())
    return sf, v.run(), v


def codes(findings):
    return Counter(f.code for f in findings)


# =============================================================================
# 一、单元测试（与垂类无关）
# =============================================================================

def test_expression():
    print("\n[单元] 表达式求值")
    nodes = {"a.x": {"value": 10}, "a.y": {"value": 3}}
    val, used = V.evaluate("a.x * a.y", nodes)
    check("乘法", val == 30, f"得到 {val}")
    check("轨迹记录引用", set(used) == {"a.x", "a.y"}, f"得到 {used}")
    check("变参 sum", V.evaluate("sum(a.x, a.y, 7)", nodes)[0] == 20)
    check("括号与小数", abs(V.evaluate("(a.x + a.y) * 0.05", nodes)[0] - 0.65) < 1e-9)

    try:
        V.evaluate("a.missing + 1", nodes)
        check("未声明节点抛 Unresolved", False, "没抛异常")
    except V.Unresolved as e:
        check("未声明节点抛 Unresolved", str(e) == "a.missing")

    for evil in ["__import__('os').system('x')", "open('/etc/passwd').read()",
                 "().__class__.__bases__[0]"]:
        try:
            V.evaluate(evil, nodes)
            check(f"拦截 {evil[:24]}", False, "没有拦截")
        except Exception:
            check(f"拦截 {evil[:24]}", True)


def test_refs():
    print("\n[单元] 引用提取")
    check("函数名不算节点", set(_refs_in_formula("sum(a.x, b.y)")) == {"a.x", "b.y"})
    check("嵌套属性链",
          set(_refs_in_formula("program.researcher.demand + 1")) == {"program.researcher.demand"})
    check("常量不算节点", _refs_in_formula("1 + 2") == [])


def test_cycles():
    print("\n[单元] 环检测")
    r2 = [{"target": "b", "formula": "c + 1"}, {"target": "c", "formula": "b + 1"}]
    check("检出二元环", len(detect_cycles(r2)) == 1 and len(detect_cycles(r2)[0]) == 3)
    r3 = [{"target": "a", "formula": "b"}, {"target": "b", "formula": "c"},
          {"target": "c", "formula": "a"}]
    check("检出三元环", len(detect_cycles(r3)) == 1)
    dag = [{"target": "b", "formula": "a + 1"}, {"target": "c", "formula": "b + 1"}]
    check("无环图不误报", detect_cycles(dag) == [])


# =============================================================================
# 二、故障注入（与垂类无关）
# =============================================================================

BASE = """
project: {id: t, name: 测试, funding_regime: 政府投资, commercial: false, physical: true}
as_of: 2026-08-06
"""


def test_schema_faults():
    print("\n[故障注入] schema 层")

    check("S1 空图判致命",
          any(f.code == "S1" and f.severity == "fatal" for f in run_all("{}")[0]))
    check("S1 缺 funding_regime",
          any(f.code == "S1" and "funding_regime" in f.subject
              for f in run_all("project: {id: t}\n")[0]))

    sf, _, _ = run_all(BASE + 'nodes:\n  a.x: {value: "一百", unit: "1"}\n')
    check("S2 字符串 value 报错", any(f.code == "S2" for f in sf),
          "这是最危险的漏检——曾静默通过")

    sf, _, _ = run_all(BASE + "nodes:\n  a.x: {value: 1}\n")
    check("S2 缺 unit 报错", any(f.code == "S2" and "unit" in f.message for f in sf))

    sf, _, _ = run_all(BASE + 'nodes:\n  a.x: {value: 1, unit: "1", provenance: E9_bogus}\n')
    check("S2 provenance 越界", any(f.code == "S2" and "provenance" in f.message for f in sf))

    sf, _, _ = run_all(BASE + '''nodes: {a.x: {value: 1, unit: "1"}}
rules:
  - {id: R1, operator: aggregate, target: b.y, formula: "a.x + z.missing", unit: "1"}
''')
    check("S3 公式引用不存在的节点", any(f.code == "S3" for f in sf))

    sf, _, _ = run_all(BASE + '''nodes: {a.x: {value: 1, unit: "1"}}
rules:
  - {id: R1, operator: bogus_op, target: b.y, formula: "a.x", unit: "1"}
''')
    check("S3 非法算子", any(f.code == "S3" and "operator" in f.message for f in sf))

    sf, _, _ = run_all(BASE + '''nodes: {a.x: {value: 1, unit: "1"}}
rules:
  - {id: R1, operator: aggregate, target: b.y, formula: "a.x +", unit: "1"}
''')
    check("S3 公式语法错误", any(f.code == "S3" and "语法" in f.message for f in sf))

    sf, _, _ = run_all(BASE + '''evidence: {EV1: {provenance: E1_given, content: x, provider: y}}
claims:
  C1: {serves: [D1], argument: {claim: x, grounds: [EV_NOPE], warrant: w, backing: [EV1]}}
''')
    check("S4 论据指向不存在的证据", any(f.code == "S4" for f in sf))

    sf, _, _ = run_all(BASE + '''sections:
  - {id: S1, title: x, covers: [项目概况]}
  - {id: S1, title: y, covers: [编制依据]}
''')
    check("S6 章节 id 重复", any(f.code == "S6" for f in sf))

    sf, _, _ = run_all(BASE + "sections:\n  - {id: S1, title: x, covers: [项目产出方桉]}\n")
    check("S5 covers 错别字", any(f.code == "S5" for f in sf),
          "错别字会让 C7 既报缺失又不认已有章节")


def test_semantic_faults():
    print("\n[故障注入] 语义层 C1-C7")

    _, f, _ = run_all(BASE + '''nodes:
  s.total: {value: 100, unit: "㎡"}
occurrences:
  s.total:
    - {section: S1, value: 100}
    - {section: S2, value: 110}
sections:
  - {id: S1, title: a, serves: [D1], covers: [项目概况]}
  - {id: S2, title: b, serves: [D1], covers: [编制依据]}
''')
    check("C1 同节点全文取值不一致", "C1" in codes(f))

    _, f, _ = run_all(BASE + '''nodes:
  a.n: {value: 10, unit: 人}
  a.i: {value: 5, unit: "㎡/人"}
  a.d: {value: 999, unit: "㎡"}
rules:
  - {id: R1, operator: index, target: a.d, formula: "a.n * a.i", unit: "㎡", basis: EV1}
evidence: {EV1: {provenance: E2_normative, content: x, published_at: 2024-01-01, retrieved_at: 2026-08-01}}
''')
    check("C2 派生值对不上", "C2" in codes(f), f"得到 {dict(codes(f))}")

    _, f, _ = run_all(BASE + '''nodes:
  a.n: {value: 10, unit: 人}
  a.i: {value: 5, unit: "㎡/人"}
  a.d: {value: 50, unit: 万元}
rules:
  - {id: R1, operator: index, target: a.d, formula: "a.n * a.i", unit: "㎡", basis: EV1}
evidence: {EV1: {provenance: E2_normative, content: x, published_at: 2024-01-01, retrieved_at: 2026-08-01}}
''')
    check("C2 单位不一致", any("单位" in x.message for x in f))

    _, f, _ = run_all(BASE + '''nodes:
  cost.total: {value: 100, unit: 万元}
  fund.a: {value: 40, unit: 万元}
  fund.b: {value: 50, unit: 万元}
  fund.total: {value: 90, unit: 万元}
rules:
  - {id: R1, operator: aggregate, target: fund.total, formula: "sum(fund.a, fund.b)",
     unit: 万元, hard_constraint: "fund.total == cost.total"}
''')
    check("C3 硬等式不成立", any(x.code == "C3" for x in f))

    _, f, _ = run_all(BASE + '''evidence: {EV1: {provenance: E1_given, content: x, provider: y}}
claims:
  C1: {serves: [D1], argument: {claim: 论点, grounds: [EV1], backing: [EV1]}}
''')
    check("C4a 断链论证（缺 warrant）", any(x.code == "C4a" for x in f))

    _, f, _ = run_all(BASE + '''evidence:
  EV1: {provenance: E5_analogical, content: x, comparability: 可比}
claims:
  C1: {serves: [D9], argument: {claim: x, grounds: [EV1], warrant: w, backing: [EV1]}}
''')
    check("C5a 类比证据独立支撑", any(x.code == "C5a" for x in f))

    _, f, _ = run_all(BASE + '''evidence:
  EV1: {provenance: E2_normative, content: x, published_at: 2018-01-01, retrieved_at: 2019-01-01}
''')
    check("C5b 证据过时效", any(x.code == "C5b" for x in f))

    _, f, _ = run_all(BASE + '''evidence:
  EV1: {provenance: E6_administrative, content: x, doc_no: A, issuer: B, approved_at: 2020-01-01}
''')
    check("C5e 批复缺有效期", any(x.code == "C5e" for x in f))

    _, f, _ = run_all(BASE + '''evidence:
  EV1: {provenance: E2_normative, content: 《普通高等学校建筑规划面积指标》建标191-2018,
        published_at: 2018-01-01, retrieved_at: 2026-08-01}
''')
    check("C5g 规范名称写法错误", any(x.code == "C5g" and x.severity == "error" for x in f),
          "编号对但名字错——真实样本里就是这么写的")

    _, f, _ = run_all(BASE + '''evidence:
  EV1: {provenance: E2_normative, content: 《普通高等学校建筑面积指标》建标191-2018,
        published_at: 2018-01-01, retrieved_at: 2026-08-01}
''')
    check("C5g 写法正确时不误报",
          not any(x.code == "C5g" and x.severity == "error" for x in f))

    _, f, _ = run_all(BASE + '''sections:
  - {id: S1, title: a, serves: [D1], covers: [项目概况], terms: {科研用房: [A, B]}}
  - {id: S2, title: b, serves: [D1], covers: [编制依据], terms: {科研用房: [A]}}
''')
    check("C6 术语跨章指称不一致", any(x.code == "C6" for x in f))

    _, f, _ = run_all(BASE + '''project: {id: t, funding_regime: 政府投资, outline_complete: true}
sections:
  - {id: S1, title: a, serves: [D1], covers: [项目概况]}
''')
    check("C7 法定内容无承载", any(x.code == "C7" for x in f))

    _, f, _ = run_all(BASE + '''sections:
  - {id: S1, title: a, serves: [D1], covers: [项目概况], content_status: missing}
''')
    check("C7b 空壳覆盖", any(x.code == "C7b" for x in f), "防止靠声明 covers 骗过 C7")

    _, f, _ = run_all(BASE + '''nodes: {a.x: {value: 1, unit: "1"}}
rules:
  - {id: R1, operator: aggregate, target: b.y, formula: "c.z + a.x", unit: "1"}
  - {id: R2, operator: aggregate, target: c.z, formula: "b.y + a.x", unit: "1"}
''')
    check("CYCLE 派生回环", any(x.code == "CYCLE" for x in f))


def test_no_false_positives():
    print("\n[故障注入] 反向：干净输入不应误报")
    _, f, v = run_all(BASE + '''nodes:
  a.n: {value: 10, unit: 人, provenance: E1_given, source: 客户}
  a.i: {value: 5,  unit: "㎡/人", provenance: E2_normative, source: 规范}
  a.d: {value: 50, unit: "㎡", provenance: E4_derived, source: 第3章}
rules:
  - {id: R1, operator: index, target: a.d, formula: "a.n * a.i", unit: "㎡", basis: EV1, serves: D4}
evidence:
  EV1: {provenance: E2_normative, content: 面积指标, published_at: 2024-01-01, retrieved_at: 2026-08-01}
''')
    check("正确派生不报 C2", "C2" not in codes(f), f"得到 {dict(codes(f))}")
    check("派生轨迹已记录", v and "a.d" in v.trace)
    check("通过项已记录", v and any("R1" in p for p in v.passed))


def test_norms_registry():
    print("\n[单元] 规范登记册")
    n = load_norms().get("norms") or {}
    check("登记册非空", len(n) >= 5)
    check("法定基准已查证", n.get("NDRC-2023-304", {}).get("status") == "verified_current")
    check("状态取值受控",
          all(v.get("status") in ("verified_current", "in_use_unverified",
                                  "superseded", "repealed", "unknown") for v in n.values()))
    check("在用未查证的条目都给了核验方法",
          all("verify" in v for v in n.values()
              if v.get("status") == "in_use_unverified" and v.get("used_by")))


# =============================================================================
# 三、垂类相关：基线全部来自 meta.yaml
# =============================================================================

def test_vertical(d, meta):
    vm = meta["vertical"]
    b = meta["test_baseline"]
    tag = vm["cn"]
    print(f"\n[垂类] {tag}（{vm['id']}）")

    for f in ["profile.yaml", "outline.yaml", "rules.yaml", "knowledge.md",
              "content-template.yaml", "fixture.yaml", "meta.yaml", "materials.yaml", "subtypes.yaml"]:
        check(f"{tag} 存在 {f}", (d / f).exists())

    prof = yaml.safe_load((d / "profile.yaml").read_text(encoding="utf-8"))["profile"]
    check(f"{tag} 继承 2023 大纲",
          prof.get("extends") in ("outline-gov-2023", "outline-ent-2023"),
          f"得到 {prof.get('extends')}")
    # 「验证过」必须说清楚验了什么、怎么验的；但**不得描述拿哪份项目验的**。
    # validated_against 记的是某个具体项目（它的口径、建设性质），那是项目信息，
    # 垂类要能用在同类型的其他项目上，不该携带它。
    val = prof.get("validation") or {}
    check(f"{tag} 声明了验证状态", "validated" in prof)
    check(f"{tag} 已验证的须写明方法与范围",
          not prof.get("validated") or (val.get("method") and val.get("verified_chains")),
          f"得到 {sorted(val)}")
    check(f"{tag} 不残留 validated_against（那是项目信息）",
          "validated_against" not in prof,
          "改用 validation.method / verified_chains 陈述规则集本身")
    check(f"{tag} meta 与 profile 的 id 一致", prof.get("id") == vm["id"])
    check(f"{tag} description 说明了不适用范围", "不适用于" in vm.get("description", ""))

    for f in ["outline.yaml", "profile.yaml", "rules.yaml", "content-template.yaml",
              "fixture.yaml", "materials.yaml", "subtypes.yaml"]:
        try:
            yaml.safe_load((d / f).read_text(encoding="utf-8")); ok, det = True, ""
        except Exception as e:
            ok, det = False, str(e)[:100]
        check(f"{tag} YAML 可解析 {f}", ok, det)

    # ---- 黄金用例 ----
    sf, findings, val = run_all((d / "fixture.yaml").read_text(encoding="utf-8"))
    n_serr = sum(1 for x in sf if x.severity in ("error", "fatal"))
    check(f"{tag} schema 无错误", n_serr == b["schema_errors"],
          f"得到 {n_serr}：{[x.subject + ':' + x.message for x in sf if x.severity != 'warning']}")
    check(f"{tag} 检查结果与基线一致", dict(codes(findings)) == b["codes"],
          f"期望 {b['codes']}\n       实际 {dict(codes(findings))}")
    check(f"{tag} 通过项 ≥ {b['min_passed']}", val and len(val.passed) >= b["min_passed"],
          f"得到 {len(val.passed) if val else 0}")
    for node in b["must_trace"]:
        check(f"{tag} 派生轨迹含 {node}", val and node in val.trace)

    t = {k: round(x["value"], 2) for k, x in val.trace.items()}
    for node, want in b["numbers"]:
        check(f"{tag} {node} = {want}", t.get(node) == want, f"得到 {t.get(node)}")

    for item in b.get("must_stay_broken") or []:
        check(f"{tag} 真实缺陷仍可见：{item['node']}",
              round(t.get(item["node"], 0)) == item["recomputed"],
              f"得到 {t.get(item['node'])}，报告声明 {item['declared_in_report']}")

    # ---- 防项目定制：结构层不得携带样本取值 ----
    import re as _re

    def numeric_leaves(o, path=""):
        if isinstance(o, dict):
            for k, v in o.items():
                yield from numeric_leaves(v, f"{path}.{k}")
        elif isinstance(o, list):
            for i, v in enumerate(o):
                yield from numeric_leaves(v, f"{path}[{i}]")
        elif isinstance(o, (int, float)) and not isinstance(o, bool):
            yield path, o

    # 允许的数值：版本号、层级、容差、比例这类结构参数
    ALLOWED = {0, 1, 2, 3, 4, 5, 6, 7, 8, 9, 10, 11,
               0.001, 0.005, 0.01, 0.1, 0.5, 0.55}
    for fn in ("profile.yaml", "outline.yaml", "rules.yaml"):
        doc = yaml.safe_load((d / fn).read_text(encoding="utf-8"))
        # 篇幅权重是结构参数（这一章占多重），不是项目取值——
        # 它对同类型的任何项目都成立，不构成「项目定制」。
        bad = [(k, v) for k, v in numeric_leaves(doc)
               if v not in ALLOWED and not k.endswith(".weight")]
        check(f"{tag} {fn} 无样本取值残留",
              not bad,
              f"发现数值型字段 {bad[:5]}——样本数字只能出现在 verified/note 这类说明文本里，"
              f"一旦成为字段取值就会被当成本项目的缺省值")

    # 操作性字段不得夹带项目专名
    # 专名清单放在 base 的 references/validation/ 里，不硬编码进代码——
    # 否则守卫本身就把要防的专名带进了交付物。生成品中该清单不存在，跳过这项。
    _np = ROOT / "references/validation/project-names.txt"
    _names = ([x.strip() for x in _np.read_text(encoding="utf-8").splitlines()
               if x.strip() and not x.startswith("#")] if _np.exists() else [])
    NAME_PAT = _re.compile("|".join(_re.escape(x) for x in _names)) if _names else None
    # 除专名外，还要挡住"描述某个具体项目"的措辞——
    # 「一份旧口径的科研用房可研（含既有建筑拆除重建）」没有专名，
    # 但它描述的是一个具体项目的口径与建设性质，同样是项目信息。
    DESC_PAT = _re.compile(r"一份(旧|新)口径|该样本(编制于|已按)|验证样本即|"
                           r"已用一份真实的[^|\n]{0,30}（")
    for fn in ("profile.yaml", "meta.yaml", "outline.yaml", "rules.yaml",
               "knowledge.md", "materials.yaml", "subtypes.yaml"):
        txt = (d / fn).read_text(encoding="utf-8")
        body = "\n".join(l for l in txt.splitlines() if not l.lstrip().startswith("#"))
        hit = DESC_PAT.findall(body)
        check(f"{tag} {fn} 无验证项目的描述性信息", not hit,
              f"命中 {set(hit)}——验证是关于规则集的陈述，不是关于某个项目的陈述")

    if NAME_PAT:
        for fn in ("outline.yaml", "content-template.yaml", "rules.yaml",
                   "profile.yaml", "knowledge.md", "materials.yaml", "meta.yaml"):
            txt = (d / fn).read_text(encoding="utf-8")
            # 注释行不算——注释是给人看的说明，不进入生成结果
            body = "\n".join(l for l in txt.splitlines() if not l.lstrip().startswith("#"))
            check(f"{tag} {fn} 正文无项目专名",
                  not NAME_PAT.search(body),
                  f"命中 {set(NAME_PAT.findall(body))}——垂类必须适用于同类型的其他项目")

    # ---- 影响分析 ----
    import impact as I
    core, argu, deriv = load_l0()
    proj = yaml.safe_load((d / "fixture.yaml").read_text(encoding="utf-8"))
    an = I.ImpactAnalyzer(proj, core, argu, deriv)
    ip = b["impact_probe"]
    down = an.downstream({ip["node"]})
    check(f"{tag} 前向传播覆盖多级下游", set(ip["expect_downstream"]) <= down,
          f"得到 {sorted(down)}")
    res = an.analyze({ip["node"]: ip["new_value"]})
    for node, exp in (ip.get("expect_delta") or {}).items():
        check(f"{tag} 派生新旧对照 {node}", res["delta"].get(node) == exp,
              f"得到 {res['delta'].get(node)}")
    check(f"{tag} 列出需重写章节", len(res["affected_sections"]) >= 1)
    allowed = set(ip.get("expect_new_error_codes") or ["C1", "C2"])
    check(f"{tag} 只报新引入的错误",
          all(e["code"] in allowed for e in res["errors_after_change"]),
          f"得到 {[e['code'] for e in res['errors_after_change']]}")

    # ---- 渲染 + 反向回读 ----
    import render as R, reconcile as RC
    content = yaml.safe_load((d / "content-template.yaml").read_text(encoding="utf-8"))
    rp = b["render_probe"]
    with tempfile.TemporaryDirectory() as td:
        rc_code = R.Renderer(proj, content).build(td)
        check(f"{tag} 渲染成功", rc_code == 0, f"退出码 {rc_code}")
        bindings = json.loads((Path(td) / "bindings.json").read_text(encoding="utf-8"))
        check(f"{tag} bindings 覆盖多处复述",
              len(bindings.get(rp["multi_occurrence_node"], [])) >= rp["min_occurrences"],
              f"得到 {len(bindings.get(rp['multi_occurrence_node'], []))} 处")
        snap = json.loads((Path(td) / "snapshot.json").read_text(encoding="utf-8"))
        check(f"{tag} 快照记录图指纹", len(snap.get("graph_sha256", "")) == 64)

        rec = RC.Reconciler(proj, bindings, core, argu, deriv)
        pb = b["reconcile_probe"]

        def occ(node, section=None):
            for o in bindings.get(node, []):
                if section is None or o["section"] == section:
                    return o
            return None

        events = []
        for key, sec_key, new_key in (("given", None, "given_new"),
                                      ("derived", "derived_section", "derived_new"),
                                      ("restated", "restated_section", "restated_new")):
            o = occ(pb[key], pb.get(sec_key) if sec_key else None)
            check(f"{tag} 回读探针节点 {pb[key]} 已渲染", o is not None)
            if o:
                events.append({"block": o["block"], "old": o["rendered"], "new": pb[new_key]})
        ch, orph, miss = rec.from_changes(events)
        r = rec.run(ch, orph, miss)
        types = {x["node"]: x["type"] for x in r["decisions"]}
        check(f"{tag} 给定事实变更 → 接受", types.get(pb["given"]) == "T_given", f"得到 {types}")
        check(f"{tag} 派生值被改 → 拒绝", types.get(pb["derived"]) == "T_derived")
        check(f"{tag} 复述处被改 → 拒绝", types.get(pb["restated"]) == "T_restated")
        check(f"{tag} 拒绝时给出上游供裁决",
              any(x.get("upstream") for x in r["decisions"] if x["type"] == "T_derived"))
        check(f"{tag} 被拒绝的不进图", set(r["accepted"]) == {pb["given"]})
        check(f"{tag} 被接受的接上影响分析", r["impact"] is not None)

        ch2, orph2, _ = rec.from_changes([{"block": "NOSUCH#9", "old": "x", "new": "y"}])
        check(f"{tag} 无绑定的块归为图外内容", len(orph2) == 1)

    bad = {"sections": [{"id": "X", "blocks": [{"type": "para", "text": "{{no.such.node}}"}]}]}
    with tempfile.TemporaryDirectory() as td:
        check(f"{tag} 未知节点中止渲染", R.Renderer(proj, bad).build(td) == 1)
        check(f"{tag} 中止时不留半成品", not (Path(td) / "report.docx").exists())

    # ---- 审计留痕 ----
    import audit as A
    ap = b["audit_probe"]
    tr = A.build_trace(proj)
    check(f"{tag} 轨迹覆盖派生节点", len(tr["derived"]) >= ap["min_derived"],
          f"得到 {len(tr['derived'])}")
    e = next((x for x in tr["derived"] if x["node"] == ap["sample_node"]), None)
    check(f"{tag} 轨迹含规则与公式", e and e["rule"] and e["formula"])
    check(f"{tag} 轨迹上溯到叶子输入", e and len(e["leaf_inputs"]) >= ap["min_leaf_inputs"])

    with tempfile.TemporaryDirectory() as td:
        A.write_trace(proj, td)
        check(f"{tag} 指纹文件已生成", (Path(td) / "trace.fingerprint.json").exists())
        check(f"{tag} 审计 verify 通过", A.verify(td) == 0)
        tp = Path(td) / "trace.json"
        tp.write_bytes(tp.read_bytes().replace(ap["tamper_token"].encode(), b"999999"))
        check(f"{tag} 篡改轨迹被检出", A.verify(td) == 1)

    with tempfile.TemporaryDirectory() as td:
        le = dict(ap["log_entry"])
        le.update({"at": "2026-08-06T10:00:00+08:00", "actor": "测试", "source": "客户编辑"})
        check(f"{tag} 写入变更日志", A.append_log(proj, dict(le), td) == 0)
        rec0 = json.loads((Path(td) / "changelog.jsonl").read_text(encoding="utf-8").splitlines()[0])
        check(f"{tag} 日志记录了影响范围", "derived_changed" in rec0["impact"])
        check(f"{tag} 链式校验通过", A.verify(td) == 0)
        lp = Path(td) / "changelog.jsonl"
        rec0["new"] = 99999
        lp.write_text(json.dumps(rec0, ensure_ascii=False, sort_keys=True) + "\n", encoding="utf-8")
        check(f"{tag} 篡改日志被检出", A.verify(td) == 1)
        check(f"{tag} 缺必填字段拒绝入账",
              A.append_log(proj, {"actor": "x", "node": "y", "old": 1, "new": 2}, td) == 2)

    # ---- intake 生成器 ----
    import intake_gen as IG
    gen = IG.IntakeGenerator(proj, core, argu, deriv)
    ask, search = gen.build()
    check(f"{tag} 生成了客户提问", len(ask) > 0)
    check(f"{tag} 按阻塞度降序",
          all(ask[i]["blocking_degree"] >= ask[i + 1]["blocking_degree"]
              for i in range(len(ask) - 1)))
    check(f"{tag} 每条提问都说明了为什么现在问", all(q.get("why") for q in ask))
    check(f"{tag} 分批提问不超额", IG.render(ask, search, batch=3).count("为什么现在问") <= 3)



# =============================================================================
# 五、素材摄入与来源路由
# =============================================================================

def test_intake(d, meta):
    """素材缺口要在动笔之前暴露，且缺什么、找谁要、不给的后果都要写出来。

    只报一句"资料不全"等于没报——客户不知道该补什么。
    """
    import intake_scan as IS
    vm = meta["vertical"]
    tag = vm["cn"]
    b = meta["test_baseline"].get("intake") or {}
    proj = yaml.safe_load((d / "fixture.yaml").read_text(encoding="utf-8"))
    mat_l0 = yaml.safe_load((ROOT / "references/materials.yaml").read_text(encoding="utf-8"))
    vert = yaml.safe_load((d / "materials.yaml").read_text(encoding="utf-8"))

    have, missing, unclassified, _src = IS.scan_materials(proj, mat_l0, vert)
    check(f"{tag} 素材台账已登记", len(have) >= b.get("min_registered_classes", 4),
          f"得到 {len(have)} 类")
    blocking = [m for m in missing if m["required"]]
    check(f"{tag} 阻塞素材数与基线一致", len(blocking) == b.get("blocking_missing", 0),
          f"期望 {b.get('blocking_missing', 0)}，得到 {len(blocking)}：{[m['class'] for m in blocking]}")
    check(f"{tag} 每条缺素材都给了追料话术与后果",
          all(m["ask"] and m["consequence"] for m in missing))
    check(f"{tag} 缺素材按阻塞在前排序",
          all(missing[i]["required"] >= missing[i + 1]["required"]
              for i in range(len(missing) - 1)))

    bad, unsourced = IS.scan_traceability(proj)
    check(f"{tag} 无坏 from（指向未登记素材或缺 locator）", not bad, str(bad)[:160])
    check(f"{tag} 无源取值数与基线一致", len(unsourced) == b.get("unsourced", len(unsourced)),
          f"期望 {b.get('unsourced')}，得到 {len(unsourced)}")

    routed, handfilled = IS.route_missing(proj, {}, vert)
    check(f"{tag} 无手填派生值", not handfilled, str(handfilled)[:160])
    check(f"{tag} 路由四类齐备", set(routed) >= {"ASK", "SEARCH", "COMPUTE", "WITHHOLD"})
    # 路由的硬规则：派生值不许出现在 ASK 清单里
    targets = {r.get("target") for r in (proj.get("rules") or []) if isinstance(r, dict)}
    check(f"{tag} ASK 清单里没有派生值",
          not (targets & {x["node"] for x in routed["ASK"]}),
          "派生值出现在 ASK 里说明路由错了，正确做法是对它的输入路由")

    # 垂类覆写必须指向合法路由
    ok_routes = {"ASK", "SEARCH", "COMPUTE", "WRITE", "WITHHOLD"}
    bad_ov = {k: v for k, v in (vert.get("routing_overrides") or {}).items()
              if v not in ok_routes}
    check(f"{tag} 路由覆写取值合法", not bad_ov, str(bad_ov))

    # ASK 回退检索：允许与禁止都要显式列出。不列 = 无从判断，agent 会自己拿主意
    elig = vert.get("ask_fallback_eligible") or []
    inel = vert.get("ask_fallback_ineligible") or []
    check(f"{tag} 声明了可回退检索的槽位", len(elig) >= 3, f"得到 {len(elig)}")
    check(f"{tag} 声明了禁止回退的槽位", len(inel) >= 3, f"得到 {len(inel)}")
    check(f"{tag} 可回退槽位都给了入图后的证据等级",
          all(e.get("provenance_after") for e in elig))
    check(f"{tag} 可回退槽位都没标成 E1_given",
          all(e.get("provenance_after") != "E1_given" for e in elig),
          "回退检索得来的值标 E1 就等于宣称它是客户给定的实测值")
    check(f"{tag} 存量类禁止回退",
          any("existing" in str(x.get("slot")) for x in inel),
          "用别处的数算缺额，缺额就是假的")
    check(f"{tag} 禁止回退的槽位都写了理由", all(x.get("why") for x in inel))

    # 篇幅权重：表达论证重心。缺了它，起草会写成均匀的流水账，
    # 重点章和陪衬章一样长，读者看不出主次。
    ol = yaml.safe_load((d / "outline.yaml").read_text(encoding="utf-8"))
    secs = [x for x in (ol.get("sections") or []) if isinstance(x, dict)]
    ws = [x for x in secs if x.get("weight") is not None]
    check(f"{tag} 章节定义了篇幅权重", len(ws) == len(secs),
          f"{len(ws)}/{len(secs)} 章有权重")
    check(f"{tag} 权重取值合理", all(0 < float(x["weight"]) <= 3 for x in ws))
    check(f"{tag} 论证落点章的权重高于平均",
          max(float(x["weight"]) for x in ws) >= 1.5,
          "每章权重都一样等于没有重心")







def test_reconcile_coverage_at_risk():
    """锚点丢失时，必须真去查法定覆盖，而不是只打印一句提示。

    这条来自一次自查：报告里写着「删除会触发 C7 法定覆盖检查」，
    但代码从没跑过那个检查。**提示一个不存在的检查比不提示更糟**——
    它让人以为已经查过了。
    """
    print("\n[回读] 删除内容的覆盖失守判定")
    import reconcile as R
    core, argu, deriv = load_l0()
    d = ROOT / "references/verticals/edu-research"
    if not (d / "fixture.yaml").exists():
        d = ROOT / "references/vertical"
    proj = yaml.safe_load((d / "fixture.yaml").read_text(encoding="utf-8"))
    secs = [x for x in (proj.get("sections") or []) if x.get("covers")]
    if not secs:
        check("垂类样例含承载法定内容的章节", False, "没有 covers 就测不了")
        return
    sec = secs[0]
    bindings = {"scale.design_total": [
        {"bookmark": "n1", "section": sec["id"], "block": "B1",
         "rendered": "13000", "value": 13000}]}
    rec = R.Reconciler(proj, bindings, core, argu, deriv)

    car = rec.coverage_at_risk(["n1"])
    check("锚点丢失会算出受影响的法定 covers", bool(car), str(car))
    check("受影响清单指到具体章节与具体 covers",
          car and car[0]["section"] == sec["id"] and car[0]["covers"] == sec["covers"])
    check("没有锚点丢失时不报", not rec.coverage_at_risk([]))
    check("丢失的锚点不在有 covers 的章节时不误报",
          not R.Reconciler(proj, {"x": [{"bookmark": "n9", "section": "CH_NONE",
                                         "rendered": "1", "value": 1}]},
                           core, argu, deriv).coverage_at_risk(["n9"]))

    out = R.render(rec.run([], [], ["n1"]))
    check("报告里列出了会失守的法定内容", "将失去承载" in out)
    check("报告区分了「整段被删」与「书签被破坏」", "书签被编辑操作破坏" in out)
    check("给了处置方向而不只是报警", "降级为节" in out)






def test_section_research():
    """章节内容项触发的检索。

    这类内容**不产生空节点**——「科研工作及学术交流」写不写、写得实不实，
    图上看不出任何缺口，缺口分诊看不见它们。不单立一条通道，
    这几章就会被写成场面话，而它们恰恰是评审判断
    「这份报告是不是认真做的」的第一印象来源。
    """
    print("\n[检索] 章节内容项触发的检索任务")
    import intake_gen as IG
    spec = yaml.safe_load(
        (ROOT / "references/section-research.yaml").read_text(encoding="utf-8"))
    tasks = spec.get("tasks") or {}
    check("定义了章节检索任务", len(tasks) >= 8)
    check("每条任务都有触发内容项", all(t.get("triggers") for t in tasks.values()))
    check("每条任务都说明产出写进哪里", all(t.get("produces") for t in tasks.values()))
    check("每条任务都定了证据级别", all(t.get("evidence") for t in tasks.values()))
    check("涉客户的任务都要求核对",
          all(tasks[k].get("must_confirm") for k in tasks if k.startswith("RS1_")
              or k.startswith("RS2_") or k.startswith("RS3_") or k.startswith("RS4_")),
          "写进单位概况的每一句都是以客户的名义在说")
    facets = spec.get("facets") or {}
    check("建设单位是必需分面", facets.get("org", {}).get("required") is True)
    check("说明了同名/过时/宣传口径三个坑",
          all(x in str(facets.get("hard_rule", "")) for x in ("同名", "过时", "宣传口径")))
    rules = spec.get("rules") or {}
    check("规定涉本单位的必须带单位名", "R2_org_facet_required" in rules)
    check("规定先核对后落笔", "R3_confirm_before_write" in rules)
    check("规定查不到不许用形容词补", "R4_no_padding" in rules)
    check("规定背景类内容要连回论点", "R5_link_to_claim" in rules)

    # 「科研工作及学术交流」这条必须存在且由「项目单位概况」触发
    rs2 = tasks.get("RS2_org_academic")
    check("含「科研工作与学术交流」任务", bool(rs2))
    check("它由项目单位概况触发", "项目单位概况" in (rs2 or {}).get("triggers", []))
    check("它的检索式带单位名分面",
          all("{org}" in q for q in (rs2 or {}).get("query", [])),
          "不带单位名查出来的是行业综述")
    check("它明确要写可核的具体项而非形容词",
          "宁可少写" in str((rs2 or {}).get("trap", "")))

    core, _argu, _deriv = load_l0()
    proj = {"project": {"owner": "某单位", "location": "某省某市", "industry": "某行业"},
            "as_of": "2026-08-06",
            "sections": [{"id": "CH1", "covers": ["项目单位概况", "编制依据"]}]}
    out = IG.research_tasks(proj, core)
    ids = {t["id"] for t in out}
    check("按骨架承载的内容项激活任务", "RS2_org_academic" in ids, str(sorted(ids)))
    check("未被骨架承载的内容项不激活", "RS12_social_risk_precedent" not in ids)
    rs = next(t for t in out if t["id"] == "RS2_org_academic")
    check("检索式已填入单位名", all("某单位" in q for q in rs["queries"]), str(rs["queries"]))
    check("须核对的排在前面", out[0]["must_confirm"] is True)

    # 缺单位名时不硬查
    proj2 = dict(proj); proj2["project"] = {"location": "某省某市"}
    out2 = IG.research_tasks(proj2, core)
    rs2b = next(t for t in out2 if t["id"] == "RS2_org_academic")
    check("缺单位全称时该任务被挡住", bool(rs2b["blocked"]), str(rs2b))
    check("挡住时说明了为什么与怎么办",
          "套话" in rs2b["blocked"] and "单位全称" in rs2b["blocked"])

    # 已完成的不重复派发
    proj3 = dict(proj); proj3["research_done"] = ["RS2_org_academic"]
    check("已完成的不再重复派发",
          "RS2_org_academic" not in {t["id"] for t in IG.research_tasks(proj3, core)})

    txt = IG.render_research(out)
    check("清单标出哪些须客户核对", "须客户核对" in txt)
    check("清单给出检不到时怎么办", "检不到" in txt)


def test_typography():
    """排版：字体字号行距页边距，并且中文字体必须单独设。

    一份内容正确但排版不合规的可研，会在收件环节就被退回——
    收件的人不看内容看格式。这是最不值得丢的分。

    重点盯 w:eastAsia：只设 font.name 的话中文会退回默认字体，
    文档里中英文两种脸，而**在装了对应字体的本机看不出来**。
    """
    print("\n[排版] 字体字号与中文字体设置")
    import render as R
    ty_doc = yaml.safe_load((ROOT / "references/typography.yaml").read_text(encoding="utf-8"))
    presets = ty_doc.get("presets") or {}
    check("定义了排版预设", len(presets) >= 2)
    check("有默认预设且存在", ty_doc.get("default_preset") in presets)
    check("字号号数对照齐全",
          {"三号", "四号", "小四", "五号"} <= set(ty_doc.get("size_map") or {}),
          "中文习惯用号数，写规格用号数、落地换算，避免「小四」被当成 12 号")
    for pn, ps in presets.items():
        el = ps.get("elements") or {}
        check(f"预设 {pn} 覆盖标题/正文/表格/图题",
              {"title", "heading1", "body", "table_cell", "caption"} <= set(el))
        check(f"预设 {pn} 正文给了行距", el["body"].get("line_spacing") is not None)
        check(f"预设 {pn} 给了页边距", bool((ps.get("page") or {}).get("margin_cm")))
    check("生僻字体都给了 fallback",
          all(("fallback" in v) or v.get("font") in ("宋体", "黑体", "楷体")
              for ps in presets.values() for v in (ps.get("elements") or {}).values()
              if v.get("font") in ("方正小标宋简体", "仿宋_GB2312", "楷体_GB2312")),
          "不是所有机器都有这些字体，缺了要退到黑体/宋体而不是方框")

    rules = ty_doc.get("rules") or {}
    check("明令中文字体必须单独设", "T1_eastasia_must_be_set" in rules)
    check("明令不嵌字体", "T2_no_font_embedding" in rules)
    check("明令排版不改变正文来源", "T3_numbers_stay_in_placeholders" in rules)
    check("明令压排版是最后手段", "T5_compact_is_last_resort" in rules)
    check("紧凑版标注了代价",
          "读得累" in str(presets.get("compact", {}).get("caution", "")))
    check("把格式要求列为必问且不可推断",
          (ty_doc.get("intake") or {}).get("cannot_infer") is True)
    check("问不到时给了可执行的替代",
          "近期通过的报告" in str((ty_doc.get("intake") or {}).get("ask_precisely", "")),
          "「请提供格式要求」客户多半没有；「给一份通过的报告」拿得到")

    # 号数换算
    sm = ty_doc["size_map"]
    check("小四 = 12 磅", R._pt({"size": "小四"}, "size", sm) == 12)
    check("三号 = 16 磅", R._pt({"size": "三号"}, "size", sm) == 16)

    # 端到端：渲染后打开成品核对
    import tempfile
    d = ROOT / "references/verticals/edu-research"
    if not (d / "fixture.yaml").exists():
        d = ROOT / "references/vertical"
    proj = yaml.safe_load((d / "fixture.yaml").read_text(encoding="utf-8"))
    content = yaml.safe_load((d / "content-template.yaml").read_text(encoding="utf-8"))
    with tempfile.TemporaryDirectory() as td:
        r = R.Renderer(proj, content)
        rc = r.build(td)
        check("带排版渲染成功", rc == 0, f"退出码 {rc}")
        docx_path = Path(td) / "report.docx"
        check("正文样式设了中文字体", R.verify_eastasia(docx_path),
              "不设的话中英文两种脸，且在本机看不出来")
        import docx as D
        from docx.oxml.ns import qn
        doc = D.Document(docx_path)
        nm = doc.styles["Normal"]
        check("正文字号已应用", nm.font.size and abs(nm.font.size.pt - 12) < .01,
              f"得到 {nm.font.size.pt if nm.font.size else None}")
        check("正文行距已应用", nm.paragraph_format.line_spacing == 1.5)
        check("首行缩进按字数算", nm.paragraph_format.first_line_indent.pt == 24,
              "中文首行缩进 2 字 = 2 × 字号，不是固定磅值")
        h1 = doc.styles["Heading 1"]._element.rPr.rFonts
        check("标题层级也设了中文字体", bool(h1.get(qn("w:eastAsia"))))
        sec = doc.sections[0]
        check("页边距已应用", abs(sec.top_margin.cm - 3.7) < .05,
              f"得到 {sec.top_margin.cm}")
        check("装订侧留得更宽", sec.left_margin.cm > sec.right_margin.cm)

        # 换一套预设应当真的换掉字号
        r2 = R.Renderer(proj, content); r2.preset = "compact"
        r2.build(td)
        d2 = D.Document(Path(td) / "report.docx")
        check("切换预设会改变正文字号",
              abs(d2.styles["Normal"].font.size.pt - 10.5) < .01,
              f"紧凑版正文应为五号 10.5，得到 {d2.styles['Normal'].font.size.pt}")


def test_length_budget():
    """篇幅：定的是权重（论证重心），不是字数。

    直接给每章定字数会诱发**凑字数**——贴台账、抄规范条文、把一句话拆成三句。
    这些都能把字数做够，每一条都在降低报告质量。
    所以权重表达的是「这一章在全篇里占多重」，字数只作为区间提示。
    """
    print("\n[篇幅] 分章权重与偏离提示")
    import render as R
    sh = yaml.safe_load((ROOT / "references/outline-shaping.yaml").read_text(encoding="utf-8"))
    lb = sh.get("length_budget") or {}
    check("L0 定义了篇幅预算", bool(lb))
    check("按报送层级给了总量档", len(lb.get("total", {}).get("defaults_by_reporting_to") or {}) >= 4)
    check("权重可按审批口重点章调整",
          any("emphasis_chapters" in str(a) for a in (lb.get("weights") or {}).get("adjust_by") or []))
    check("用区间而非定值", "区间而非定值" in str((lb.get("bands") or {}).get("cn", "")))
    hr = str(lb.get("hard_rules") or [])
    check("明令字数不足不许注水", "不许靠注水补" in hr)
    check("明令超长先移附件不删内容", "先移附件" in hr)
    check("明令篇幅不构成删法定内容的理由", "从不构成删除法定内容的理由" in hr)
    # 从「提示」改成「门禁」是有意的：提示没有强制力，某章被一笔带过，
    # 报告照样出得来，然后在评审那里被指出来。
    g = lb.get("gate") or {}
    check("篇幅核对是门禁", "不放行" in str(g.get("rule")), str(g)[:120])
    check("门禁给出缺口清单而不只是判定", "缺口清单" in str(g.get("but")))
    check("门禁的三条出口含书面下调目标",
          any("下调" in str(x) for x in (g.get("exits") or [])), str(g.get("exits")))
    check("门禁仍禁止注水", "充数" in str(g.get("forbidden")))
    pct = lb.get("per_chapter_target") or {}
    check("分章字数是客户输入且优先于权重", "优先于权重" in str(pct.get("priority")))
    check("表格不计入字数", "表格与插图不计" in str(pct.get("note")))

    cp = sh.get("client_profile") or {}
    check("把「哪几章看得最细」列为必问", "emphasis_chapters" in cp)
    pcl = cp.get("per_chapter_length") or {}
    check("逐章字数要求单独问", bool(pcl.get("question")), str(cp.keys()))
    check("逐章字数问不到时接受部分登记", "部分登记优于不登记" in str(pcl.get("note")))
    check("逐章字数要在起草前问", "起草前" in str(pcl.get("order")))
    check("重点章标为不可推断", cp["emphasis_chapters"].get("cannot_infer") is True,
          "只有做过这个审批口的人才知道，agent 查不到")

    # 字数统计
    check("中文按字计", R._cjk_len("总建筑面积一万三千平方米") == 12,
          f"得到 {R._cjk_len('总建筑面积一万三千平方米')}")
    check("英文数字折半计", R._cjk_len("PUE1.25") == 3, f"得到 {R._cjk_len('PUE1.25')}")

    # 偏离判定
    import tempfile
    with tempfile.TemporaryDirectory() as td:
        op = Path(td) / "outline.yaml"
        op.write_text(yaml.safe_dump(
            {"sections": [{"id": "CH1", "weight": 1.0}, {"id": "CH3", "weight": 2.0}]},
            allow_unicode=True), encoding="utf-8")
        txt, off = R.length_report({"CH1": 300, "CH3": 600}, str(op))
        check("按权重分配后均衡的不报偏离", not off, str(off))
        txt2, off2 = R.length_report({"CH1": 900, "CH3": 100}, str(op))
        kinds = {c: k for c, _n, _t, k in off2}
        check("重点章被一笔带过会报偏短", kinds.get("CH3") == "偏短",
              "让「某章被一笔带过」在出稿前可见，而不是等评审指出来")
        check("次要章写过头会报偏长", kinds.get("CH1") == "偏长")
        check("偏离时给出处理方向", "补事实" in txt2 and "移附件" in txt2)
    txt3, off3 = R.length_report({"CH1": 300}, None)
    check("没有权重时只报字数不判偏离", not off3 and "只报字数" in txt3)



def test_design_depth():
    """可研阶段的设计深度：agent 做到哪一步。

    早先这里一刀切成「设计一律不做」，那是错的。可研里的设计是**方案深度**，
    其中规范判定与参数推导 agent 该做也能做；只有方案比选与设计单位专属的
    才需要外部输入。混成一句「去问客户」，会把该自己算的推给客户，
    也会让该设计单位定的被 agent 编掉。
    """
    print("\n[设计] 可研阶段的设计深度分层")
    dp = yaml.safe_load((ROOT / "references/discipline-probes.yaml").read_text(encoding="utf-8"))
    dd = dp.get("design_depth") or {}
    check("定义了设计深度分层", bool(dd))
    layers = {x["layer"] for x in (dd.get("agent_does") or [])}
    check("agent 承担规范判定", "norm_lookup" in layers)
    check("agent 承担参数推导", "computed" in layers)
    check("agent 给方案候选但不替客户选", "draft_with_candidates" in layers)
    nd = {x["layer"] for x in (dd.get("agent_does_not") or [])}
    check("设计单位专属的仍不做", "design_only" in nd)
    check("规范判定必须给条款依据",
          any("条款依据" in str(x.get("rule", "")) for x in (dd.get("agent_does") or [])),
          "只给结论不给依据，评审会当成拍脑袋，而且事后无从复核")
    check("给了可研深度的判据", "depth_test" in dd)

    probes = dp.get("common_probes") or []
    OK_RES = {"norm_lookup", "computed", "choice", "design_only", "mixed", "ask"}
    bad = [p["id"] for p in probes if p.get("resolution") not in OK_RES]
    check("通用取值点都标了 resolution", not bad, str(bad))
    need_by = [p["id"] for p in probes
               if p.get("resolution") not in ("ask",) and not p.get("resolve_by")]
    check("非 ask 的取值点都说明了怎么定", not need_by, str(need_by))

    # 参数推导模板
    deriv = yaml.safe_load((ROOT / "references/derivations.yaml").read_text(encoding="utf-8"))
    t = deriv.get("templates") or {}
    for tid in ("T18_load_from_index", "T19_transformer_capacity", "T20_hvac_load",
                "T21_water_demand", "T22_fire_water_volume", "T23_fire_compartment"):
        check(f"含设计推导模板 {tid.split('_')[0]}", tid in t)
    check("冷热负荷模板点明分母是空调面积",
          "空调面积不是总建筑面积" in str(t["T20_hvac_load"].get("caution", "")),
          "用总建筑面积算会系统性偏大")
    check("用水量模板点明人数口径要一致",
          "口径" in str(t["T21_water_demand"].get("caution", "")))
    check("防火分区模板只定上限不定划分",
          "只定" in str(t["T23_fire_compartment"].get("caution", "")))

    # 分组输出
    import claim_derive as CD
    dpl = dp
    sub = {"subtypes": {}}
    proj = {"project": {"id": "t"}, "nodes": {}, "rules": []}
    ps = CD.probe_status(CD.collect_probes(dpl, sub, {}), proj)
    groups = {p.get("resolution") for p in ps}
    check("取值点能按 resolution 分组", len(groups) >= 3, str(groups))


def test_section_prompts():
    """逐章起草提示：每一章是不同的活，提示词就该不同。

    用一套通用提示词写十几章，出来的必然是同一种腔调的流水账——
    每章都七百字，读者看不出主次，评审看不出哪一章是这个项目的重点。
    """
    print("\n[起草] 逐章提示与切片绑定")
    import claim_derive as CD
    doc = yaml.safe_load((ROOT / "references/section-prompts.yaml").read_text(encoding="utf-8"))
    secs = doc.get("sections") or {}
    check("逐章定义了起草提示", len(secs) >= 8, f"得到 {len(secs)}")
    for k, v in secs.items():
        check(f"{k} 绑定了切片", bool(v.get("slice")))
        check(f"{k} 说明了核心问题", bool(v.get("focus")))
        check(f"{k} 列了这一章特有的坑", len(v.get("must_not") or []) >= 2)
    check("每章的坑各不相同",
          len({tuple(v.get("must_not") or []) for v in secs.values()}) == len(secs),
          "坑一样说明提示词没有真正分章")

    cm = doc.get("common") or {}
    check("通用约束不在各章重复", cm.get("hard") and cm.get("style"))
    check("技术方案章按设计深度分层准备",
          all(x in str(secs["technical"].get("prepare")) for x in
              ("norm_lookup", "computed", "choice", "design_only")),
          "这一章的准备工作正是设计深度那四层")
    check("结论章禁止出现新数",
          "新数" in str(secs["conclusion"].get("must_not")))
    check("规模章点名口径混用",
          "口径混用" in str(secs["scale"].get("must_not")))
    check("必要性章点名断链",
          "断链" in str(secs["background"].get("must_not")))
    check("起草有前置门禁", "C18_prepare_before_draft" in (doc.get("constraints") or {}))

    p = CD.section_prompt("technical")
    check("能取出某章的提示", p and p["section"].get("cn"))
    txt = CD.render_prompt("technical", p)
    check("提示里含起草前置", "起草前必须先完成" in txt)
    check("提示里含通用约束", "通用约束" in txt)
    check("未知章节给出可用清单", "section-prompts" in CD.render_prompt("nope", None))


def test_figures():
    """插图生成：数据图可以画，设计图一律不画。

    这条界线不是风格问题。数据图错了顶多是数错；设计图编了，
    是把一份不存在的专业判断伪装成设计成果——评审会当成设计单位的活来审，
    审出来的问题记到设计单位头上。
    """
    print("\n[插图] 数据图生成与设计图禁令")
    import figures as F
    spec = F.load_spec()
    cat = spec.get("figures") or {}
    forb = spec.get("forbidden") or []
    check("定义了可生成的图", len(cat) >= 6)
    check("定义了禁止生成的图", len([x for x in forb if isinstance(x, dict) and x.get("kind")]) >= 5)
    kinds = " ".join(str(x.get("kind")) for x in forb if isinstance(x, dict))
    for must in ("总平面图", "结构布置图", "效果图"):
        check(f"禁止生成：{must}", must in kinds)
    check("工艺流程图也在禁令内", "工艺流程图" in kinds,
          "它看起来像数据图，其实出自工艺设计——这一类最容易被误判")
    check("每条禁令都给了替代做法",
          all(x.get("instead") for x in forb if isinstance(x, dict) and x.get("kind")))
    rules = spec.get("rules") or {}
    check("规定图上每个数都来自节点", "R1_data_only" in rules)
    check("规定缺节点直接中止", "R2_abort_on_gap" in rules)
    check("规定必须标为示意图", "R3_label_as_schematic" in rules)
    check("规定正式图到位即让位", "R5_yield_to_real" in rules)

    d = ROOT / "references/verticals/edu-research"
    if not (d / "fixture.yaml").exists():
        d = ROOT / "references/vertical"
    proj = yaml.safe_load((d / "fixture.yaml").read_text(encoding="utf-8"))
    if "cost.construction" in (proj.get("nodes") or {}):
        svg = F.build("F_COST_COMPOSITION", proj, spec)
        check("能生成投资构成图", svg.startswith("<svg") and svg.endswith("</svg>"))
        check("图题标明是示意图", "示意图" in svg)
        check("图注列出了数据来源节点", "cost.construction" in svg)
        nums = [str(v.get("value")) for k, v in proj["nodes"].items()
                if k.startswith("cost.") and isinstance(v.get("value"), (int, float))]
        check("图上的数确实来自节点", any(n in svg for n in nums))

        # 分项之和与合计不符 → 中止，不出图
        bad = yaml.safe_load(yaml.safe_dump(proj))
        bad["nodes"]["cost.total"] = {"value": 99999, "unit": "万元", "provenance": "E4_derived"}
        try:
            F.build("F_COST_COMPOSITION", bad, spec)
            ok = False
        except F.FigureError as e:
            ok = "不符" in str(e)
        check("分项之和对不上合计会中止", ok,
              "「图漏了一个分项」由此变成可机检的问题，而不是靠人看出来")

    # 缺节点 → 中止
    try:
        F.build("F_SCHEDULE", {"nodes": {}}, spec)
        ok = False
    except F.FigureError as e:
        ok = "不要自己拆" in str(e) or "schedule.milestones" in str(e)
    check("缺里程碑时不自行拆工期", ok, "拆出来的节点会被当成承诺")

    # 设计图 → 拒绝，并给替代做法
    try:
        F.build("F_SITE_PLAN", proj, spec)
        ok, msg = False, ""
    except F.FigureError as e:
        msg = str(e); ok = True
    check("请求设计图会被拒绝", ok)
    check("拒绝时列出禁令类型并给替代做法",
          "总平面图" in msg and "设计单位" in msg, msg[:120])

    # 渲染集成
    r = (ROOT / "scripts/render.py").read_text(encoding="utf-8")
    check("render 支持 figure 块", 'elif t == "figure"' in r)
    check("插图失败会进 errors 而不是静默跳过", "self.errors.append(f\"插图" in r,
          "静默跳过会让一份少了图的稿子看起来像是完整的")


def test_deliverable_is_docx():
    """交付物必须是 Word 文档，不是 project.yaml。

    这条测试来自一次真实的失败：SKILL.md 的「第四步：出稿」里只写了
    validators.py 和 audit.py，**没有 render.py**。agent 照着做，
    最后交出去的是一个 YAML 文件。图是脚手架，不是产品。
    """
    print("\n[交付] 成品是 docx，不是 YAML")
    md = (ROOT / "SKILL.md").read_text(encoding="utf-8")
    check("description 里说明了交付 Word 文档",
          "Word 文档" in md.split("---")[1],
          "触发描述不提交付物，agent 会以为产物是图")
    check("正文有独立的交付物说明", "交付物" in md)
    check("明确 project.yaml 不单独交付",
          "不单独交付" in md or "中间产物" in md)
    check("出稿步骤调用了 render.py",
          "scripts/render.py project.yaml" in md,
          "只写 validators.py 和 audit.py 的话，agent 永远走不到生成 docx 那一步")
    check("说明了用 SendUserFile 交付", "SendUserFile" in md)
    check("有交付前自查清单", "交付前自查" in md)
    check("自查里点名了 report.docx", "report.docx" in md)

    # 渲染器确实产出 docx，且路径就是文档里写的那个
    r = (ROOT / "scripts/render.py").read_text(encoding="utf-8")
    check("render.py 输出 report.docx", 'out / "report.docx"' in r)
    check("render.py 输出回读所需的 bindings", "bindings.json" in r)



def test_ask_fallback():
    """ASK 回退检索：客户给不出时，检索候选请其选。

    这条路径与编造只隔一层标注。锁住的就是那层标注：
    候选要客户选、入图不得升级为 E1_given、必须写明何时替换成实定值。
    """
    print("\n[回退] ASK 回退检索与数值候选")
    import fit_rank as FR
    routing = yaml.safe_load(
        (ROOT / "references/provenance-routing.yaml").read_text(encoding="utf-8"))
    fb = routing.get("ask_fallback") or {}
    check("定义了 ASK 回退检索", bool(fb))
    check("列出了允许回退的槽位类型", len(fb.get("eligible", {}).get("classes") or []) >= 4)
    check("列出了禁止回退的槽位类型", len(fb.get("ineligible", {}).get("classes") or []) >= 5)
    inelig = str(fb.get("ineligible") or {})
    for must in ("现状描述", "存量台账", "人员与办学规模", "行政文件"):
        check(f"禁止回退：{must}", must in inelig)
    pr = fb.get("provenance_rules") or {}
    check("明令不得升级为 E1_given", "永远不升级为 E1_given" in str(pr.get("never_upgrade", "")))
    check("参照值有生命周期要求", "不允许静默保留" in str(pr.get("placeholder_lifecycle", "")))
    check("规定了正文措辞", "参照值" in str(pr.get("in_text_wording", "")))
    bd = str(fb.get("boundary_with_NF1") or {})
    check("与 NF1 的界线给了正反例", "违反 NF1" in bd and "合规" in bd)

    spec = yaml.safe_load((ROOT / "references/external-fit.yaml").read_text(encoding="utf-8"))
    vfd = spec.get("value_fit_dimensions") or {}
    check("数值候选有独立的评分维度", len(vfd) >= 6)
    weights = [v["weight"] for v in vfd.values() if isinstance(v, dict) and "weight" in v]
    check("口径匹配权重最高",
          vfd["N1_metric_definition"]["weight"] == max(weights),
          "口径对不上时报告读起来完全正常，所以它最危险")
    check("数值候选不默认勾选",
          all(b.get("default_checked") is False
              for b in (vfd.get("ranking") or {}).get("buckets", {}).values()),
          "默认勾选会让客户顺手确认一个他没细看的数")
    check("候选表必须留两个出口", len(vfd.get("must_offer") or []) == 2)

    proj = {"project": {"location": "某省某市"}}
    cands = [
        {"id": "V1", "value": 65, "unit": "kWh/(㎡·a)", "source_doc": "某标准",
         "authority": "强制性标准", "value_type": "约束值",
         "metric_definition_fit": 3, "object_fit": 2, "geography_fit": 2, "vintage_fit": 3},
        {"id": "V2", "value": 82, "unit": "kWh/(㎡·a)", "source_doc": "某统计报告",
         "authority": "地方主管部门", "value_type": "统计平均值",
         "metric_definition_fit": 1, "object_fit": 3, "geography_fit": 3, "vintage_fit": 2},
        {"id": "V3", "value": 55, "source_doc": "某媒体报道", "authority": "媒体报道"},
    ]
    kept, dropped = FR.rank_values(proj, cands)
    check("二手转述被剔除", [d["id"] for d in dropped] == ["V3"], str(dropped))
    check("口径匹配的排第一", kept[0]["id"] == "V1", str([k["id"] for k in kept]))
    check("口径存疑的排后但不剔除", any(k["id"] == "V2" for k in kept))
    check("没有一条默认勾选", not any(k["default_checked"] for k in kept))
    check("每条都注明入图后标什么等级",
          all(k["provenance_after"] != "E1_given" for k in kept))
    check("口径未写明时显式提示",
          "未写明口径" in FR.rank_values(
              proj, [{"id": "X", "value": 1, "authority": "国家标准"}])[0][0]["metric_definition"])

    import validators as V
    core, argu, deriv = load_l0()

    def c14(node):
        p2 = {"project": {"id": "t"}, "nodes": {"energy.index": node}, "rules": []}
        v = V.Validator(core, argu, deriv, p2)
        v.c14_placeholder()
        return [f for f in v.findings if f.code == "C14"]

    check("参照值标成 E1_given 会报错",
          any(f.severity == "error" for f in c14(
              {"value": 65, "provenance": "E1_given", "is_placeholder": True,
               "adopted_by": "x", "refine_when": "y"})),
          "标 E1 等于宣称这是客户给定的实测值，责任被悄悄记到客户头上")
    check("参照值缺 adopted_by 会报错",
          any(f.severity == "error" for f in c14(
              {"value": 65, "provenance": "E2_normative", "is_placeholder": True,
               "refine_when": "y"})))
    check("参照值缺 refine_when 会警告",
          any(f.severity == "warning" for f in c14(
              {"value": 65, "provenance": "E2_normative", "is_placeholder": True,
               "adopted_by": "x"})))
    check("标注齐全的参照值只出提示不报错",
          all(f.severity == "info" for f in c14(
              {"value": 65, "provenance": "E2_normative", "is_placeholder": True,
               "adopted_by": "x", "refine_when": "y"})))
    check("E5 类比值被当取值会报错",
          any(f.severity == "error" for f in c14({"value": 6400, "provenance": "E5_analogical"})),
          "同类项目的值只能佐证合理性，不能充当取值")
    check("显式标 benchmark_only 的 E5 不报错",
          not c14({"value": 6400, "provenance": "E5_analogical", "benchmark_only": True}))


def test_shaping():
    """骨架变形：形可以随甲方变，神不能变。

    「神」的机检定义就是 covers 全集不变。变形只搬动内容的容器，
    不增减内容本身——这条不变量守不住，三层解耦就白做了。
    """
    print("\n[骨架] 变形算子与不变量")
    doc = yaml.safe_load((ROOT / "references/outline-shaping.yaml").read_text(encoding="utf-8"))
    ops = doc.get("operators") or {}
    check("定义了变形算子", len(ops) >= 6)
    check("每个算子都说明了对 covers 的影响",
          all("covers_effect" in v for k, v in ops.items() if k != "forbidden"),
          str([k for k, v in ops.items() if k != "forbidden" and "covers_effect" not in v]))
    forb = str((ops.get("forbidden") or {}).get("list") or [])
    check("明令禁止删除法定 covers", "删除法定 covers" in forb)
    check("明令禁止一个 cover 由多章承载", "多个章承载" in forb)
    check("明令禁止 G1 后改章号", "G1 冻结之后" in forb)

    inv = doc.get("invariants") or {}
    check("covers 守恒是第一不变量", "I1_covers_conserved" in inv)
    check("covers 唯一承载有独立不变量", "I1b_covers_unique" in inv)
    check("章节顺序受数据依赖约束", "I2_dependency_order" in inv)

    cp = doc.get("client_profile") or {}
    check("甲方基本情况覆盖报送层级与审批惯例",
          "reporting_to" in cp and "reviewer_convention" in cp)
    check("报送层级标为不可推断", cp["reporting_to"].get("cannot_infer") is True)
    check("审批惯例标为不可推断", cp["reviewer_convention"].get("cannot_infer") is True,
          "地方审批口的惯例成文规范里查不到，只能问客户")

    rules = doc.get("shaping_rules") or []
    check("变形规则给了触发条件与理由",
          all(r.get("when") and r.get("suggest") and r.get("rationale") for r in rules))
    sr8 = next((r for r in rules if r["id"] == "SR8"), None)
    check("篇幅压力不构成删 covers 的理由",
          sr8 and "篇幅压力不构成删 covers 的理由" in str(sr8.get("hard_rule", "")))


def test_claim_derive(d, meta):
    """论证推演与专业取值点。

    这两件事都不靠垂类手写死——子类是长尾的，穷举不了也维护不动。
    锁住的是：识别得出子类、取值点能按子类激活、缺论据能分派路由、
    以及**识别不出子类时不许退回通用条款**。
    """
    import claim_derive as CD
    vm = meta["vertical"]
    tag = vm["cn"]
    b = meta["test_baseline"].get("derive") or {}
    proj = yaml.safe_load((d / "fixture.yaml").read_text(encoding="utf-8"))
    sub = yaml.safe_load((d / "subtypes.yaml").read_text(encoding="utf-8"))
    dp = yaml.safe_load((ROOT / "references/discipline-probes.yaml").read_text(encoding="utf-8"))
    cd = yaml.safe_load((ROOT / "references/claim-derivation.yaml").read_text(encoding="utf-8"))
    _c, argu, _dv = load_l0()

    subs = sub.get("subtypes") or {}
    check(f"{tag} 定义了子类", len(subs) >= b.get("min_subtypes", 3), f"得到 {len(subs)}")
    check(f"{tag} 每个子类都有识别线索与评审关注点",
          all(v.get("recognize") and v.get("review_focus") for v in subs.values()))
    check(f"{tag} 每个子类都显式声明了可写通用条款的专业",
          all("generic_clause_ok" in v for v in subs.values()),
          "不列 = 不允许写通用条款，所以这个字段必须显式出现")
    all_probes = [p for v in subs.values() for p in (v.get("probes") or [])]
    check(f"{tag} 子类取值点都问到了具体内容",
          all(p.get("ask") and "？" in str(p["ask"]) for p in all_probes))
    common_ids = {p["id"] for p in (dp.get("common_probes") or [])}
    dup = [p["id"] for p in all_probes if p["id"] in common_ids]
    check(f"{tag} 子类不重复通用取值点", not dup, f"重复 {dup}")
    good_disc = set(dp.get("disciplines") or {})
    baddisc = [p["id"] for p in all_probes if p.get("discipline") not in good_disc]
    check(f"{tag} 取值点挂在合法专业下", not baddisc, str(baddisc))

    hits, _s, _n = CD.recognize_subtypes(proj, sub)
    check(f"{tag} 能从项目图与名称识别出子类", len(hits) >= 1, str(hits))
    probes = CD.probe_status(CD.collect_probes(dp, sub, hits), proj)
    check(f"{tag} 取值点按子类激活（多于通用集）",
          len(probes) > len(common_ids), f"得到 {len(probes)}")
    check(f"{tag} 已答与已确认不涉及的不再列为缺口",
          any(p["status"] == "satisfied" for p in probes)
          and any(p["status"] == "confirmed_na" for p in probes))

    claims, extra = CD.derive_claims(cd, argu, proj, sub, hits)
    check(f"{tag} 推出了论证需求", len(claims) >= 8, f"得到 {len(claims)}")
    check(f"{tag} 缺失论据都带了来源路由",
          all(g.get("route") for c in claims for g in c["missing_grounds"]))
    check(f"{tag} 带出了垂类额外论点",
          any(e.get("source") == "vertical" for e in extra), str([e.get("source") for e in extra]))
    # 已声明 grounds_patterns 的维度不该再被报缺同一模式
    d2 = next((c for c in claims if c["dimension"] == "D2"), None)
    check(f"{tag} 已落地的论据模式不再重复报缺",
          d2 and "G6_norm_compliance" not in [g["pattern"] for g in d2["missing_grounds"]])

    # 未识别子类时的兜底：取并集反问，不退回通用条款
    blank = {"project": {"name": "某项目", "profile": vm["id"]}}
    hits0, _s0, _n0 = CD.recognize_subtypes(blank, sub)
    check(f"{tag} 无线索时不硬猜子类", not hits0, str(hits0))
    union = CD.collect_probes(dp, sub, list(subs))
    check(f"{tag} 兜底取并集覆盖全部子类取值点",
          len(union) >= len(common_ids) + len({p["id"] for p in all_probes}))

def test_intake_sla():
    """追料时限。没有时限的等待会一直等下去——直到有人随手把那个数填上。

    这里锁三件事：
      ① 时钟按工作日走，且「正在办理中」不重置它（否则永远不到期）
      ② 到期只有四条合法出口，「估一个」不在其中
      ③ 到决策点必须主动暴露，不能等客户来问
    """
    print("\n[素材] 追料时限与到期处置")
    import datetime as _dt
    import intake_scan as IS
    mat_l0 = yaml.safe_load((ROOT / "references/materials.yaml").read_text(encoding="utf-8"))
    sla = mat_l0.get("intake_sla") or {}
    check("L0 定义了追料时限", bool(sla.get("profiles")))
    check("时限分档递增（催办 < 升级 < 决策）",
          all(p["remind_days"] < p["escalate_days"] < p["decide_days"]
              for p in sla["profiles"].values()))
    exits = sla.get("exits") or {}
    check("到期有四条合法出口",
          {"S1_substitute", "S2_withhold_and_risk", "S3_narrow_scope", "S4_suspend"} <= set(exits))
    forbidden = str((exits.get("forbidden") or {}).get("list") or [])
    check("到期明确禁止估算与套用",
          "估一个数" in forbidden and "套用同类项目" in forbidden and "经验值" in forbidden)
    check("每个素材类都挂了时限档",
          all("sla" in v for k, v in (mat_l0.get("material_classes") or {}).items()
              if k not in ("M8_norms", "M0_unclassified")),
          str([k for k, v in (mat_l0.get("material_classes") or {}).items() if "sla" not in v]))

    # 工作日计算：跨周末不能按自然日算
    d0, d1 = _dt.date(2026, 7, 1), _dt.date(2026, 7, 15)   # 14 自然日
    check("按工作日计龄（跨周末不虚增）", IS._bdays(d0, d1) == 10,
          f"得到 {IS._bdays(d0, d1)}")

    def probe(asked, cls="M4_cost_finance", note=None, as_of=_dt.date(2026, 8, 6)):
        g = {"class": cls, "status": "pending", "asked_at": asked}
        if note:
            g["note"] = note
        return IS.scan_aging({"as_of": as_of, "material_gaps": [g]}, mat_l0)[0]

    check("新发出的追料不打扰", probe(_dt.date(2026, 8, 5))["stage"] == "open")
    check("到点催办", probe(_dt.date(2026, 7, 29))["stage"] == "remind",
          probe(_dt.date(2026, 7, 29))["stage"])
    check("到点升级", probe(_dt.date(2026, 7, 20))["stage"] == "escalate",
          probe(_dt.date(2026, 7, 20))["stage"])
    check("到点决策", probe(_dt.date(2026, 6, 20))["stage"] == "decide")

    # 「正在办理中」不重置时钟——这是最常见的规避方式
    stalled = probe(_dt.date(2026, 6, 20), note="客户答复正在办理中")
    check("「正在办理中」不重置时钟", stalled["stage"] == "decide",
          "答一句「在办」就能无限延期的话，这套时限等于没有")

    # 行政批复类给更长的时限，且默认出口是 S2 而非死等
    adm = probe(_dt.date(2026, 7, 20), cls="M1_approval")
    check("行政批复类时限更宽松", adm["stage"] == "remind", adm["stage"])
    check("行政批复类默认出口是留白进风险章",
          "S2" in str((sla["profiles"]["administrative"] or {}).get("note", "")))

    # 没有 asked_at 等于没追过
    noclock = IS.scan_aging(
        {"as_of": _dt.date(2026, 8, 6),
         "material_gaps": [{"class": "M4_cost_finance", "status": "pending"}]}, mat_l0)[0]
    check("没有 asked_at 会被点名", noclock["stage"] == "unknown")

    # 已解决的不再计时
    done = IS.scan_aging(
        {"as_of": _dt.date(2026, 8, 6),
         "material_gaps": [{"class": "M4_cost_finance", "status": "resolved",
                            "asked_at": _dt.date(2026, 6, 1)}]}, mat_l0)
    check("已补齐的不再计时", not done)

    # 到决策点必须排在最前，主动暴露
    mixed = IS.scan_aging({"as_of": _dt.date(2026, 8, 6), "material_gaps": [
        {"class": "M6_baseline", "status": "pending", "asked_at": _dt.date(2026, 8, 5)},
        {"class": "M4_cost_finance", "status": "pending", "asked_at": _dt.date(2026, 6, 20)},
    ]}, mat_l0)
    check("到期项排在最前（主动暴露，不等客户问）", mixed[0]["stage"] == "decide")

def test_intake_without_vertical():
    """没有匹配垂类时，素材扫描不得报"缺素材：无"。

    实测踩过：base 拿一份非高校项目去扫，只给了一份设计说明，
    输出是"缺素材：无。本垂类要求的素材类均已登记"——
    一个假的全绿，读起来像资料齐了。**比不检查更危险。**
    修法是 L0 给出缺省必备素材集作下限，并在报告里写明比对依据是哪一套。
    """
    print("\n[素材] 无垂类兜底")
    import intake_scan as IS
    mat_l0 = yaml.safe_load((ROOT / "references/materials.yaml").read_text(encoding="utf-8"))
    check("L0 有缺省必备素材集", bool(mat_l0.get("default_required_materials")))

    proj = {"project": {"id": "t", "name": "某市政道路改造工程", "nature": "改建",
                        "funding_regime": "政府投资", "location": "某省某市",
                        "physical": True},
            "materials": [{"id": "M3-01", "class": "M3_design", "name": "方案设计说明"}],
            "nodes": {"scale.length": {"value": 3200, "unit": "m",
                                       "provenance": "E1_given", "source": "设计方案"}},
            "rules": []}
    have, missing, unc, src = IS.scan_materials(proj, mat_l0, {})
    blocking = [m for m in missing if m["required"]]
    check("无垂类时仍报出缺素材", len(missing) > 0,
          "报「缺素材：无」是假的全绿——比不检查更危险")
    check("无垂类时仍识别出阻塞类", len(blocking) >= 3, f"得到 {len(blocking)}")
    check("显式说明比对依据是缺省集", "缺省" in src, src)
    check("每条缺素材都有追料话术与后果",
          all(m["ask"] and m["consequence"] for m in missing))

    # 非实体工程不该被要求用地与设计素材
    proj2 = dict(proj)
    proj2["project"] = dict(proj["project"], physical=False)
    _, missing2, _, _ = IS.scan_materials(proj2, mat_l0, {})
    blocked2 = {m["class"] for m in missing2 if m["required"]}
    check("纯软件/纯服务项目不强制用地与设计素材",
          "M2_planning_land" not in blocked2 and "M3_design" not in blocked2,
          str(sorted(blocked2)))

    # 缺省集只是下限：有垂类时必须让位给垂类
    vert = yaml.safe_load(
        (ROOT / "references/verticals/edu-research/materials.yaml").read_text(encoding="utf-8")
    ) if (ROOT / "references/verticals/edu-research/materials.yaml").exists() else None
    if vert is None:
        vert = yaml.safe_load(
            (ROOT / "references/vertical/materials.yaml").read_text(encoding="utf-8"))
    _, _, _, src2 = IS.scan_materials(proj, mat_l0, vert)
    check("有垂类时以垂类要求为准", "缺省" not in src2, src2)

def test_fit_rank():
    """外部证据不得直接入图：跨地域的必须被硬否决，采纳与否交客户。"""
    print("\n[外部证据] 适配性评估与排序")
    import fit_rank as FR
    proj = {"project": {"location": "某省某市", "funding_regime": "政府投资",
                        "nature": "新建", "profile": "demo"}, "as_of": None}
    cands = [
        {"id": "A", "title": "本市配建规定", "authority": "地方规范性文件",
         "status": "现行有效", "verified_at": None, "jurisdiction": ["某省某市"],
         "subjects": ["不限"], "clause": "第五条", "quantified": True,
         "category_fit": 3, "proposed_role": ["benchmark"]},
        {"id": "B", "title": "邻省配建规定", "authority": "地方规范性文件",
         "status": "现行有效", "jurisdiction": ["他省"], "subjects": ["不限"],
         "clause": "第三条", "quantified": True, "category_fit": 3},
        {"id": "C", "title": "已废止的旧办法", "authority": "部门规章",
         "status": "已废止", "jurisdiction": ["全国"], "subjects": ["不限"]},
        {"id": "D", "title": "仅限企业投资项目的办法", "authority": "部门规章",
         "status": "现行有效", "jurisdiction": ["全国"], "subjects": ["企业投资"]},
        {"id": "E", "title": "某行业发展白皮书", "authority": "白皮书",
         "status": "现行有效", "jurisdiction": ["全国"], "subjects": ["不限"],
         "principle_only": True, "category_fit": 1},
    ]
    kept, dropped, over = FR.rank(proj, cands, top=6)
    vetoed = {d["id"]: d["veto"] for d in dropped}
    check("跨地域的被硬否决", vetoed.get("B") == "V2_jurisdiction_miss", str(vetoed))
    check("已废止的被硬否决", vetoed.get("C") == "V1_repealed", str(vetoed))
    check("主体不适用的被硬否决", vetoed.get("D") == "V4_subject_miss", str(vetoed))
    ids = [k["id"] for k in kept]
    check("同级地方规定排第一", ids and ids[0] == "A", str(ids))
    top = next(k for k in kept if k["id"] == "A")
    check("建议采纳的默认勾选", top["default_checked"] and top["bucket"] == "recommended",
          str(top["bucket"]))
    bg = next(k for k in kept if k["id"] == "E")
    check("白皮书只能作背景且不默认勾选",
          bg["bucket"] == "background" and not bg["default_checked"], str(bg["bucket"]))
    check("每条都说明了为什么排在这个位置", all(k["why"] for k in kept))
    check("硬否决项留痕（不是静默丢弃）", len(dropped) == 3, str(len(dropped)))

    # 超限要显式说明，不许静默截断
    many = [dict(cands[0], id=f"X{i}") for i in range(9)]
    _, _, over2 = FR.rank(proj, many, top=6)
    check("超出呈报上限时显式报数", over2 == 3, f"得到 {over2}")

# =============================================================================
# 四、L0 资产自检
# =============================================================================

def _mini_docx(path):
    """造一份最小历史报告：两章、带下级标题、正文里埋着各类**必须被脱掉**的东西。"""
    import docx
    d = docx.Document()
    d.add_heading("项目建设内容及规模", 1)
    d.add_paragraph("本项目由示范大学建设，位于示范市示范路。")
    d.add_heading("科研用房", 2)
    d.add_paragraph("总建筑面积 18319 平方米，2019 年开工。")
    d.add_paragraph("依据示范发改〔2019〕第 88 号批复及 GB 50352-2019 执行。")
    d.add_paragraph("张三教授牵头，李四、王五、赵六等参与。")
    d.add_heading("配套用房", 2)
    d.add_paragraph("配套用房 2000 平方米。")
    t = d.add_table(rows=2, cols=3)
    for i, c in enumerate(["指标名称", "数量", "单位"]):
        t.rows[0].cells[i].text = c
    d.add_heading("投资估算及资金筹措", 1)
    d.add_paragraph("综上，投资估算合理，项目是可行的。")
    d.save(str(path))


def test_corpus():
    print("\n[语料库] 入库脱值 / 切章 / 蒸馏")
    import corpus as C
    tmp = Path(tempfile.mkdtemp())
    docx_p = tmp / "hist.docx"
    try:
        _mini_docx(docx_p)
    except Exception as e:
        check("python-docx 可用", False, str(e)[:80])
        return
    graph = {"nodes": {"总建筑面积": {"value": 18319}}, "sections": []}
    meta = {
        "id": "T-1", "names": ["示范大学"], "places": ["示范市", "示范路"],
        "orgs": ["科研用房"],
        "fit_meta": {"vertical": "edu-research", "subtypes": ["ST_DRY_LAB"],
                     "reporting_to": "省级发改", "outline_version": "2023",
                     "year": 2019, "review_outcome": "通过"},
        "section_map": {"^项目建设内容": "scale", "^投资估算": "investment"},
    }
    e = C.ingest(str(docx_p), graph, meta, str(tmp / "T-1"))

    check("切出两章", set(e["sections"]) == {"scale", "investment"}, str(list(e["sections"])))
    # 正文现在在 chunks/*.md 里，投影断言直接读 MD——顺带验证落盘内容也是干净的
    body = "\n".join(f.read_text(encoding="utf-8")
                     for f in sorted((tmp / "T-1" / "chunks").glob("scale-*.md")))
    check("数值已投影为节点名", "⟨总建筑面积⟩" in body, body)
    check("对不上节点的数值降级为 ⟨数值⟩", "⟨数值⟩" in body)
    check("建设单位已投影", "示范大学" not in body)
    check("地名已投影", "示范市" not in body and "示范路" not in body)
    check("文号已投影", "⟨文号⟩" in body and "88 号" not in body)
    check("规范编号已投影", "⟨规范编号⟩" in body and "50352" not in body)
    check("年份已投影", "⟨年份⟩" in body)
    check("人名已投影", "张三" not in body and "李四" not in body, body)
    # 这是本文件最要紧的一条：投影漏一个数字，参照时就会把它带进新报告
    check("投影自检无残留数值", not e["projection_leftovers"],
          str(e["projection_leftovers"]))

    # 正文落 chunks/*.md，entry.yaml 只留形——检索时物理上带不出原文
    check("entry.yaml 不再存正文", "paragraphs" not in e["sections"]["scale"],
          str(list(e["sections"]["scale"])))
    cdir = tmp / "T-1" / "chunks"
    check("每个 chunk 一个 MD", cdir.exists() and any(cdir.glob("*.md")))
    idx = yaml.safe_load((tmp / "T-1" / "index.yaml").read_text(encoding="utf-8"))
    check("有 chunk 索引", bool(idx.get("chunks")), str(idx)[:120])
    ids = [c["chunk"] for c in idx["chunks"]]
    check("chunk_id 用章节key-序号，不用哈希",
          all(re.match(r"^[a-z_]+-\d{2}$", i) for i in ids), str(ids))
    md = (cdir / f"{ids[0]}.md").read_text(encoding="utf-8")
    check("MD 里写明不得进起草上下文", "不要放进起草上下文" in md)
    check("MD 里是脱值骨架", "18319" not in md and "示范大学" not in md)

    # ---- 投影撞号：标错的节点名比不标更有害 --------------------------------
    # 表格暴露过一次：冬季温度 18℃ 被标成 ⟨fire.roof_tank_volume⟩（水箱 18 m³）、
    # 噪声 50dB 标成 ⟨storm.overflow_period⟩（重现期 50 年）。参照者会当真。
    pj = C.build_projector(
        {"nodes": {"a.tank": {"value": 18}, "b.temp": {"value": 18},
                   "c.area": {"value": 18319}, "d.occ": {"value": 1.5},
                   "e.floors": {"value": 12}}}, {})
    check("小整数不映射（必然撞号）", "12" not in pj["val2node"], str(pj["val2node"]))
    check("同值多节点不映射", "18" not in pj["val2node"], str(pj["val2node"]))
    check("大数值映射到节点", pj["val2node"].get("18319") == "c.area", str(pj["val2node"]))
    check("带小数的值映射到节点", pj["val2node"].get("1.5") == "d.occ", str(pj["val2node"]))
    # 表格单元格里开头的「1.」是数值的一部分，不是列表序号。
    # 不区分会把「1.5 ㎡/人」切成「1.」+「5」，投出「1.⟨某节点⟩」——比不投影还糟。
    check("单元格开头的数不当序号剥", C.project("1.5", pj, True) == "⟨d.occ⟩",
          C.project("1.5", pj, True))
    check("正文里开头的序号仍当序号", C.project("1. 设计依据", pj).startswith("1."))

    sec = e["sections"]["scale"]
    check("捕到表格栏目", any(t["cols"][:1] == ["指标名称"] for t in sec["tables"]),
          str(sec["tables"]))
    # 只捕表头不够：设计参数大半在表里，正文里的数字反而多是规范限值
    check("捕到表体", any(t.get("body") for t in sec["tables"]), str(sec["tables"])[:200])
    check("捕到下级小节", len(sec["subs"]) == 2, str(len(sec["subs"])))
    # 切 chunk 与排逐段计划必须用同一个函数，否则「参照这一节」和
    # 「写这一节」指的不是一回事
    import section_brief as _SB
    check("切 chunk 与逐段计划共用分组逻辑",
          _SB._rollup.__doc__ and "同一个函数" in _SB._rollup.__doc__)
    sh = sec["shape"]
    check("蒸馏出段落功能序列", bool(sh["flow"]), str(sh))
    check("蒸馏出成分配比", bool(sh["mix"]))
    check("蒸馏结果里没有原文", "paragraphs" not in sh)
    check("结论章以结论判断收尾",
          e["sections"]["investment"]["shape"]["ends_with"] == "结论判断",
          str(e["sections"]["investment"]["shape"]))

    # 切章质量：有标题样式时不该报"没有标题样式"
    warns = " ".join(e["split_quality"]["warnings"])
    check("有标题样式时不误报", "没有标题样式" not in warns, warns)

    # ---- 引用完整性：语料与图谱之间唯一的那根线 --------------------------
    r = C.ref_integrity(str(tmp / "T-1"), graph)
    check("能核出文本锚点", r["refs"] >= 1, str(r)[:150])
    check("无悬空标签时不报错", not r["dangling"], str(r["dangling"]))
    # 改个节点名，语料里的标签就该被报出来——松耦合不等于不核
    r2 = C.ref_integrity(str(tmp / "T-1"), {"nodes": {"改了名的.节点": {"value": 1}}})
    check("图谱改名后报悬空标签", bool(r2["dangling"]), str(r2)[:150])
    check("覆盖率如实计算", 0 <= r["coverage"] <= 1, str(r["coverage"]))
    # 类型槽位（⟨数值⟩⟨文号⟩）不是节点引用，不该混进来
    check("类型槽位不计入节点引用",
          not (set(C.slot_refs(str(tmp / "T-1"))) & C.TYPED_SLOTS),
          str(sorted(C.slot_refs(str(tmp / "T-1")))))

    hits = C.rank_entries([e], "scale", {"vertical": "edu-research",
                                         "subtype": "ST_DRY_LAB",
                                         "reporting_to": "省级发改"})
    check("同垂类同子类同口径命中", len(hits) == 1 and hits[0][0] >= 9, str(hits[:1]))
    txt = C.render_retrieve("scale", hits, {"vertical": "edu-research"})
    # 检索产物必须是提示词。骨架一旦进起草上下文，模型就改写而不是重写
    check("默认不吐骨架原文", "⟨总建筑面积⟩" not in txt)
    check("提示词含铺陈顺序", "铺陈顺序" in txt)
    check("提示词含表格栏目", "栏目" in txt)
    txt2 = C.render_retrieve("scale", hits, {"vertical": "edu-research"},
                             show_skeleton=True)
    check("显式索取才给骨架", "骨架" in txt2)

    d2 = C.diff_hints({"reporting_to": "省级发改", "outline_version": "2023前旧口径",
                       "subtypes": ["ST_DRY_LAB"], "known_gaps": ["需求分析"]},
                      {"reporting_to": "报部委", "outline_version": "2023",
                       "subtype": "ST_WET_LAB"})
    check("差异提示报口径不同", any("口径" in x for x in d2), str(d2))
    check("旧大纲提示结构不可参照", any("结构不可参照" in x for x in d2), str(d2))
    check("缺失章节提示不要继承", any("不要连缺失一起继承" in x for x in d2), str(d2))

    check("空库时不报错", "没有可比条目" in C.render_retrieve("scale", [], {}))


def test_section_brief():
    print("\n[逐章简报] 查 + 借鉴 + 逐段计划")
    import section_brief as SB
    # 小节数落在可排表区间时按小节；层级过细时卷到最粗的一层
    subs = [{"title": f"{p}设计", "level": 2, "paras": 5, "chars": 500,
             "mix": ["定量陈述 3", "定性说明 2"]} for p in "建结水电暖"]
    deep = []
    for s in subs:
        deep.append(s)
        deep += [{"title": f"{s['title']}-{i}", "level": 3, "paras": 3,
                  "chars": 200, "mix": ["定量陈述 3"]} for i in range(4)]
    rows = SB._rollup(deep)
    check("按最粗层级卷起", len(rows) == 5, str(len(rows) if rows else None))
    check("孙节并入父节字数", rows[0]["chars"] == 500 + 800, str(rows[0]))
    check("列出下含小节", len(rows[0]["sub_titles"]) == 4, str(rows[0].get("sub_titles")))

    flat = [{"title": f"节{i}", "level": 2, "paras": 2, "chars": 100, "mix": []}
            for i in range(40)]
    check("小节过多时不卷（排不出表）", SB._rollup(flat) is None)

    shape = {"subs": subs, "char_count": 2500, "para_count": 25,
             "slots": [{"func": "定量陈述", "paras": 25, "chars": 2500}]}
    spec = {"section": {"covers": ["技术方案"], "must_state": ["结构设计取值"]}}
    plan = SB.paragraph_plan(shape, spec)
    check("排出逐段计划", plan["rows"] is not None and len(plan["rows"]) == 5)
    check("总字数守恒", abs(plan["total"] - 2500) <= 5 * 80, str(plan["total"]))
    # 挂不上的必写项不许硬塞进第一行——那是把漏写藏起来
    check("挂不上的必写项如实列为跨单元",
          "技术方案" in plan["unplaced"], str(plan))
    txt = SB.render_plan(plan, "technical")
    check("计划里写明一段一交", "一段一交" in txt)
    check("计划里允许增删段", "段数可以不等于历史" in txt)
    check("跨单元项要逐条核销", "逐条核销" in txt)

    # 预算优先于历史篇幅：客户给了字数要求就按客户的来
    plan2 = SB.paragraph_plan(shape, spec, budget_chars=5000)
    check("有预算时按预算缩放", plan2["total"] > plan["total"] * 1.5,
          f"{plan['total']} → {plan2['total']}")

    # 无语料库时不能罢工，要退到通用段型
    plan3 = SB.paragraph_plan(None, spec)
    check("无历史样本仍出计划", plan3["rows"] is not None)
    check("无历史样本时如实标注",
          "通用段型" in SB.render_plan(plan3, "technical"))

    for f in ["_budget", "corpus_block", "research_for_section", "build"]:
        check(f"section_brief 暴露 {f}", hasattr(SB, f))


def test_length_gate():
    print("\n[篇幅门禁] 分章字数是输入，不达标不出稿")
    import length_gate as LG
    from assets import load_all as _la
    core, _argu, _deriv, _norms = _la()

    proj = {
        "project": {"vertical": "edu-research"},
        "nodes": {"a.x": {"value": 1}, "a.y": {"value": 2},
                  "a.z": {"value": None}},
        "sections": [
            {"id": "CH1", "length_target": 1000, "covers": ["项目概况"],
             "provides": ["a.x", "a.y", "a.z"]},
            {"id": "CH2", "length_target": 100, "covers": ["项目建设背景"]},
            {"id": "CH3", "covers": ["需求分析"]},
        ],
    }
    content = {"sections": [
        {"id": "CH1", "blocks": [
            {"type": "para", "text": "本项目面积{{a.x}}平方米。"},
            {"type": "table", "title": "一张很长的表" * 200},
        ]},
        {"id": "CH2", "blocks": [{"type": "para", "text": "背" * 200}]},
        {"id": "CH3", "blocks": [{"type": "para", "text": "无目标章" * 5}]},
    ]}

    counts = LG.chapter_counts(content)
    # 表格不计入字数——否则贴一张大表就能把门禁蒙过去
    check("表格不计入字数", counts["CH1"] < 60, str(counts))
    check("正文计入字数", counts["CH2"] >= 190, str(counts))

    tgt, src = LG.targets(proj, None, None, counts)
    check("客户指定的目标被采纳", tgt.get("CH1") == 1000 and src["CH1"] == "客户指定")
    check("未指定的章不判定", "CH3" not in tgt, str(tgt))

    rows = LG.gate(proj, content, core)
    v = {r[0]: r[3] for r in rows}
    check("差得远的判未达标", v["CH1"] == "未达标", str(v))
    check("超上限的判超上限", v["CH2"] == "超上限", str(v))
    check("未设目标的不判定", v["CH3"] == "未设目标", str(v))

    gaps = [r for r in rows if r[0] == "CH1"][0][5]
    titles = " ".join(g[0] for g in gaps)
    # 缺口清单是这个脚本的主产出。只判定不给清单，补字数就只剩把句子写长
    check("报出已有值但未引用的节点", "没引用的节点" in titles, titles)
    used = [g for g in gaps if "没引用的节点" in g[0]][0][1]
    check("已引用的节点不算缺口", "a.x" not in used and "a.y" in used, str(used))
    check("空值节点单列一类", "仍为空的节点" in titles, titles)
    empt = [g for g in gaps if "仍为空" in g[0]][0][1]
    check("空值节点识别正确", empt == ["a.z"], str(empt))

    txt = LG.render_gate(rows)
    check("门禁给出三条合法出口", "补事实" in txt and "补论证" in txt and "下调" in txt)
    check("门禁明示禁止注水", "抄规范条文充数" in txt)
    check("超上限先移附件不删内容", "不删内容" in txt, txt[-400:])

    # 章号（CH3）与起草提示 key（scale）按 covers 交集认亲
    spec = LG._prompt_spec("CH3", {"covers": ["需求分析", "建设内容和规模"]})
    check("章号能对上起草提示", bool(spec.get("must_state")), str(spec)[:120])

    ok = LG.gate({"project": {}, "nodes": {}, "sections": [
        {"id": "CH1", "length_target": 10}]},
        {"sections": [{"id": "CH1", "blocks": [
            {"type": "para", "text": "十二个字的一段话正好"}]}]}, core)
    check("达标的章放行", ok[0][3] == "达标", str(ok))
    check("全达标时明说可以出稿", "可以出稿" in LG.render_gate(ok))


def test_volatility():
    print("\n[时效性] 软约束可被新证据替换，硬约束不行")
    import yaml as _y
    from assets import REF
    v = _y.safe_load((REF / "volatility.yaml").read_text(encoding="utf-8")) or {}
    vc = v.get("volatile_classes") or {}
    check("软约束分类齐备", set(vc) >= {f"V{i}" + s for i, s in [
        (1, "_norm_version"), (2, "_policy_document"), (3, "_index_quota"),
        (4, "_price_cost"), (5, "_statistics"), (6, "_benchmark"),
        (7, "_org_public_info")]}, str(sorted(vc)))
    for k, spec in vc.items():
        check(f"{k} 有保鲜期", isinstance(spec.get("ttl_days"), int), str(spec)[:80])
        check(f"{k} 说明谁能替换", bool(spec.get("replace_by")), str(spec)[:80])
    # 换版最要紧的一条：只换编号不换条款号比不换更危险
    check("规范换版要复核条款号", "条款号" in str(vc["V1_norm_version"]))
    # 单位公开信息不接受检索单方替换——写进单位概况的每句都以客户名义在说
    check("单位公开信息只能由客户素材替换",
          vc["V7_org_public_info"]["replace_by"] == ["客户素材"],
          str(vc["V7_org_public_info"]["replace_by"]))
    check("单位公开信息不接受 SEARCH 单方替换",
          "不接受 SEARCH 单方替换" in str(vc["V7_org_public_info"].get("hard")))

    r = v.get("replacement") or {}
    # 这是本文件的核心：把软约束当冲突抛给客户，客户也不知道选哪个，
    # 事情就一直挂着，最后带着旧版本出稿
    check("软约束替换不算冲突", "不走客户裁决" in str(r.get("not_a_conflict")))
    check("替换仍要过适配性评估", "external-fit" in str(r.get("still_gated")))
    check("替换要留痕", "变更记录" in str(r.get("record")))
    check("定额造价替换后必须重算", "impact.py" in str(r.get("recompute")))
    check("语料库里的规范默认已过期", "默认已过期" in str(r.get("from_corpus")))

    hc = v.get("hard_constraints") or {}
    check("硬约束不许被检索替换", "不许被检索结果替换" in str(hc.get("rule")))
    check("硬约束边界情形已交代", "客户素材里的定额引用" in str(hc.get("boundary_case")))

    c = v.get("constraints") or {}
    check("有软约束保鲜检查 C21", "C21_volatile_verified" in c)
    check("只查被引用的软约束", "只查**被引用的**" in str(c["C21_volatile_verified"]))
    check("有硬约束防自动替换检查", "C21c_hard_not_auto_replaced" in c)

    # 各专业设计依据清单：形（该引哪些）可复用，值（版本号）不可复用
    n = _y.safe_load((REF / "norms.yaml").read_text(encoding="utf-8")) or {}
    db = n.get("design_basis_by_discipline") or {}
    disc = db.get("disciplines") or {}
    check("设计依据清单按专业分组", len(disc) >= 6, str(sorted(disc)))
    check("清单条目数达量级", sum(len(x) for x in disc.values()) >= 50,
          str(sum(len(x) for x in disc.values())))
    check("只记 code_seen 不记权威编号",
          all("code_seen" in it and "code" not in it
              for v2 in disc.values() for it in v2))
    check("明令 code_seen 不得直接抄进报告", "抄进报告" in str(db.get("usage", {}).get("never")))
    check("全部条目视同未核验", "不声称任何一条现行有效" in str(db.get("status_of_all_entries")))
    check("地标不入清单只留槽位", bool(db.get("local_slots")))


def test_literal_numbers():
    print("\n[字面数字] 表格里手打的数字与图无关，改图不会跟着改")
    from render import literal_numbers
    # 该报的：表格单元格里手打的事实
    check("手打的面积被报出", literal_numbers("18319") == ["18319"])
    check("手打的金额被报出", literal_numbers("总投资 13032 万元") == ["13032"])
    # 不该报的：这些数字不是事实
    for t, why in [("第5章", "章节号"), ("1.2 设计依据", "小节号"), ("（3）本工程", "序号"),
                   ("2019年", "年份"), ("GB50189-2015", "规范编号"),
                   ("〔2023〕304号", "文号"), ("m3/d", "单位指数")]:
        check(f"不误报{why}：{t}", not literal_numbers(t), str(literal_numbers(t)))
    # {{节点}} 解析之后才渲染，占位符本身不含字面数字
    check("占位符不算字面数字", not literal_numbers("{{scale.design_total}}"))

    # 表的行也该受 SSOT 管：图上加一个分项，表自动多一行
    import render as RD
    r = RD.Renderer({"nodes": {"a.x.area": {"value": 1, "unit": "㎡", "note": "甲"},
                               "a.y.area": {"value": 2, "unit": "㎡", "note": "乙"},
                               "b.z.area": {"value": 3, "unit": "㎡"}}}, {})
    rows = r._table_rows({"rows_from": {"prefix": "a.", "suffix": ".area"}})
    check("按节点族生成表行", len(rows) == 2, str(rows))
    check("行值写成占位符而非字面值",
          all(x[1].startswith("{{") for x in rows), str(rows))
    check("用 note 当行标签", rows[0][0] == "甲", str(rows))
    r2 = RD.Renderer({"nodes": {}}, {})
    r2._table_rows({"rows_from": {"prefix": "zzz.", "suffix": ".area"}})
    check("空表要报错不能静默出稿", bool(r2.errors), str(r2.errors))


def test_assets():
    print("\n[资产] L0 完整性")
    ref = ROOT / "references"
    for f in ["ontology-core.yaml", "ontology-argumentation.yaml", "outline-gov-2023.yaml",
              "outline-ent-2023.yaml", "triage.yaml", "interaction.yaml",
              "derivations.yaml", "reconcile.yaml", "norms.yaml",
              "materials.yaml", "provenance-routing.yaml", "external-fit.yaml",
              "outline-shaping.yaml", "discipline-probes.yaml", "claim-derivation.yaml",
              "figures.yaml", "typography.yaml", "section-research.yaml",
              "section-prompts.yaml", "corpus.yaml", "volatility.yaml"]:
        check(f"存在 {f}", (ref / f).exists())

    for f in sorted(ref.rglob("*.yaml")):
        try:
            yaml.safe_load(f.read_text(encoding="utf-8")); ok, det = True, ""
        except Exception as e:
            ok, det = False, str(e)[:100]
        check(f"YAML 可解析 {f.relative_to(ROOT)}", ok, det)

    core, argu, _ = load_l0()
    check("法定锚点存在", "regulatory_anchor" in core)
    check("法定锚点为 304 号文",
          "304" in str((core.get("regulatory_anchor") or {}).get("authority", {}).get("doc", "")))
    dims = (argu.get("feasibility_frame") or {}).get("dimensions") or {}
    check("维度为 D1-D11", set(dims) == {f"D{i}" for i in range(1, 12)}, f"得到 {sorted(dims)}")


# =============================================================================

def main():
    vs = verticals()
    mode = "垂类" if (ROOT / "references/vertical/meta.yaml").exists() else "base"
    print("=" * 70)
    print(f"可行性研究报告 skill —— 测试套件（{mode} 模式，{len(vs)} 个垂类）")
    print("=" * 70)

    test_expression()
    test_refs()
    test_cycles()
    test_schema_faults()
    test_semantic_faults()
    test_no_false_positives()
    test_norms_registry()
    test_deliverable_is_docx()
    test_figures()
    test_design_depth()
    test_section_prompts()
    test_length_budget()
    test_typography()
    test_section_research()
    test_reconcile_coverage_at_risk()
    test_shaping()
    test_ask_fallback()
    test_intake_sla()
    test_intake_without_vertical()
    test_fit_rank()
    test_corpus()
    test_length_gate()
    test_volatility()
    test_literal_numbers()
    test_section_brief()
    for d, meta in vs:
        test_vertical(d, meta)
        test_intake(d, meta)
        test_claim_derive(d, meta)
    test_assets()

    print("\n" + "=" * 70)
    total = _passed + len(_failed)
    if _failed:
        print(f"失败 {len(_failed)} / {total}")
        for n, _ in _failed:
            print(f"  · {n}")
        return 1
    print(f"全部通过：{_passed} / {total}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
